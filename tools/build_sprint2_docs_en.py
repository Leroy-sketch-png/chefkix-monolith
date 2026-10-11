"""Build the English, Times New Roman Sprint 2 documentation edition."""

from pathlib import Path
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_sprint2_docs import draw_flow

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "sprints" / "2026-10-11"
EVIDENCE = ROOT / "output" / "sprints" / "evidence"
OUT.mkdir(parents=True, exist_ok=True)
INK = RGBColor(28, 40, 52)
BLUE = RGBColor(35, 69, 91)


def table(doc, headers, rows, keep_together=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    available = 16.0
    widths = [available / len(headers)] * len(headers)
    if len(headers) == 2:
        widths = [4.5, 11.5]
    elif len(headers) == 3:
        widths = [2.8, 5.9, 7.3]
    for row_index, row in enumerate([headers, *rows]):
        cells = t.rows[0].cells if row_index == 0 else t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].width = Cm(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cells[i]._tc.get_or_add_tcPr()
            mar = tcpr.find(qn("w:tcMar"))
            if mar is None:
                mar = OxmlElement("w:tcMar")
                tcpr.append(mar)
            for side, value_twips in (("top", 130), ("bottom", 130), ("left", 150), ("right", 150)):
                node = OxmlElement("w:" + side)
                node.set(qn("w:w"), str(value_twips))
                node.set(qn("w:type"), "dxa")
                mar.append(node)
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8EEF2")
                tcpr.append(shade)
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.12
                if row_index == 0 or (keep_together and row_index < len(rows)):
                    para.paragraph_format.keep_with_next = True
                for run in para.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    run.font.bold = row_index == 0
        trpr = t.rows[row_index]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            rep = OxmlElement("w:tblHeader")
            rep.set(qn("w:val"), "true")
            trpr.append(rep)
    return t


def base(title, subtitle):
    d = Document()
    s = d.sections[0]
    s.top_margin = Cm(1.8)
    s.bottom_margin = Cm(1.7)
    s.left_margin = Cm(2.3)
    s.right_margin = Cm(2.3)
    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
        st = d.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0) if name == "Title" else BLUE
        st.font.bold = name != "Title"
        st.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        st.paragraph_format.space_after = Pt(5)
    # Built-in Word heading styles retain theme-font attributes unless removed.
    # Make the requested typeface explicit across paragraph and character styles.
    for st in d.styles:
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for key in list(fonts.attrib):
            if key.rsplit("}", 1)[-1].lower().endswith("theme"):
                del fonts.attrib[key]
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + key), "Times New Roman")
    tpr = d.styles["Title"].element.get_or_add_pPr()
    for old in tpr.findall(qn("w:pBdr")):
        tpr.remove(old)
    bd = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "nil")
    bd.append(bot)
    tpr.append(bd)
    p = d.add_paragraph(style="Title")
    p.add_run(title)
    p = d.add_paragraph(subtitle)
    p.runs[0].font.name = "Times New Roman"
    p.runs[0].font.size = Pt(11)
    p.runs[0].italic = True
    p = d.add_paragraph("IRON CHEF v3  |  Vo Trung Tin  |  Reviewed 11 October 2026")
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = BLUE
    f = s.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = f.add_run("IRON CHEF  |  Sprint 2  |  Page ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    f._p.append(field)
    return d


def h(doc, title, level=1):
    doc.add_heading(title, level)


def p(doc, value):
    doc.add_paragraph(value)


def bullets(doc, values):
    for value in values:
        x = doc.add_paragraph(value, style="List Bullet")
        x.paragraph_format.space_after = Pt(3)


def numbered(doc, values):
    for i, value in enumerate(values, 1):
        x = doc.add_paragraph(f"{i}. {value}")
        x.paragraph_format.space_after = Pt(3)


def sources(doc):
    h(doc, "Source baseline and interpretation")
    table(doc, ["Source", "Evidence used"], [
        ("Sprint 2 report supplied by the user", "Work assignment: Vo Trung Tin owns UI requirements, use cases, data flows, and Member scaffold acceptance. The sprint runs 23 September–10 October 2026."),
        ("Member backlog, branch codex/member-sprint12-deliverables", "Epic 1–12 implementation states and open gates G1–G10; engineering checkpoint dated 6 October 2026."),
        ("Lead backlog v109, origin/master 8bb8167", "As of 9 October 2026, Stage 3 is artifact-verified through exact step 327/561. Segment 6 was launched toward step 374. Final model quality and scoring remain unverified."),
        ("Browser captures dated 6 October 2026", "Actual frontend views of /explore/graph and /scan, with the graph backed by local seed data."),
        ("Frontend and monolith code", "Routes, request boundaries, error states, fields, and adapter behavior reviewed for this edition."),
    ])
    p(doc, "Evidence rule: a browser image proves that a screen appeared during that run; code shows an implemented path; a backlog records the team's checkpoint. None alone proves production deployment, model quality, clinical safety, or a research result.")


def figure(doc, file, caption):
    doc.add_picture(str(EVIDENCE / file), width=Cm(12.5))
    x = doc.add_paragraph(caption)
    x.alignment = WD_ALIGN_PARAGRAPH.CENTER
    x.paragraph_format.space_after = Pt(8)
    for r in x.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(9)
        r.italic = True


def build_srs():
    d = base("Sprint 2 Software Requirements Specification", "Member interface requirements and Lead integration boundaries")
    p(d, "This specification covers Vo Trung Tin's Sprint 1–2 deliverables. The graph and scan screens have run in a real browser. The graph capture contains 25 seeded ingredients and four substitution relationships; it does not show a completed Lead graph. The photo screen shows controls, not an accepted detector. Requirements below therefore separate implemented interfaces from acceptance gates.")
    table(d, ["Document control", "Value"], [
        ("Edition", "English revision 1.2, 11 October 2026"),
        ("Audience", "Member and Lead developers, reviewers, testers, and thesis supervisors"),
        ("Scope", "User, Creator, and reviewer UI; graph, scan, substitution, allergens, feedback, evaluation, and thesis evidence"),
        ("Outside Member acceptance", "Canonical datasets, model training and scoring, reviewed clinical policy, production deployment, and human research"),
    ])
    sources(d)
    h(d, "1 Product context and actors")
    p(d, "ChefKix combines recipe creation and cooking journeys with ingredient intelligence. The Lead owns scientific evidence, canonical data, models, policies, and evaluation protocols. The Member owns user-facing surfaces and adapters. A visible UI control must not be treated as proof that its external model or evidence source is accepted.")
    table(d, ["Actor", "Goal and access boundary"], [
        ("Guest", "May view public welcome, graph, and scan screens. Account-dependent actions may stop at sign-in or integration-pending states."),
        ("User", "May save allergen preferences, inspect recipes, scan ingredients, select substitutions, cook, and submit feedback under an authenticated profile."),
        ("Creator", "May draft, edit, and publish recipes while preserving ingredient quantities and recipe context."),
        ("Reviewer or admin", "May inspect protected evaluation and thesis evidence views with provenance and claim limits."),
        ("Lead or AI provider", "Supplies accepted model artifacts, manifests, graph data, and reviewed policy. UI code cannot independently certify these inputs."),
    ])
    h(d, "2 Observed browser interface")
    figure(d, "sprint2-03-graph-redesigned-top.jpg", "Figure 1. Graph search, relationship filters, coverage summary, and detail panel in the browser on 6 October 2026.")
    p(d, "The running graph reported 25 ingredients, seven with links, and four documented substitution relationships. It loaded through the configured Knowledge API, but this was a local seed collection. Chemical-similarity and co-occurrence filters are visible controls; their presence does not establish those data types in the seed response.")
    figure(d, "sprint2-03-graph-search-centered.jpg", "Figure 2. Searching Ginger centers an isolated node; compound and nutrition evidence is unavailable.")
    figure(d, "sprint2-02b-scan-controls.jpg", "Figure 3. Public scan screen with camera, image upload, and dish-photo controls. No live detector inference is evidenced.")
    h(d, "3 Functional requirements")
    requirements = [
        ("UI-01 Recipe authoring", "The Creator can draft, edit, and publish recipes. Ingredient identities and quantities remain available as context for substitutions.", "Existing workflow; verify the saved context across services."),
        ("UI-02 Scan and input review", "The User can capture or upload an image and see distinct permission, loading, error, unavailable, and demo states. Detections must be reviewable and editable before downstream use; manual input must remain possible.", "Camera, upload, and overlay exist. Per-detection editing and real detector/device acceptance remain open."),
        ("UI-03 Contextual substitutions", "Send ingredient, reason, recipe context, dietary tags, saved allergen flags, and session identity. Show candidate, quantity ratio, confidence, source, and unavailable values separately.", "Adapters and UI exist; verify live source and receipt chain."),
        ("UI-04 Allergen profile", "Support the configured EU 14 and FDA 9 groups plus custom terms; persist normalized flags and forward saved values on later requests.", "Settings and persistence exist; authenticated end-to-end verification remains open."),
        ("UI-05 Three-valued safety", "Safe requires reviewed policy support. Missing, invalid, conflicting, or unknown evidence yields Check. A policy conflict yields Blocked and disables the primary action.", "Resolver and regression tests exist; production policy acceptance is pending."),
        ("UI-06 Compound explanation", "Show source, grounding, units, and unavailable measurements. Compound presence or overlap must not imply cooking suitability, sensory equivalence, nutrition equivalence, or allergen safety.", "UI surface exists; official index and real-payload capture are pending."),
        ("UI-07 Knowledge graph", "Search and center results, drag and pan or zoom, inspect nodes and edges, load bounded neighborhoods, and show actual coverage.", "UI and bounded seed API work; official graph import remains open."),
        ("UI-08 Evaluation dashboard", "Protect the reviewer route, distinguish complete, pending, and rejected results, hide pending numbers, and preserve provenance and claim limits in exports.", "Dashboard and schema validation exist; verified Lead manifests and scores are pending."),
        ("UI-09 Feedback", "Record accept, reject, skip, and post-cooking outcomes against a session and real candidate receipt; show failure and retry states.", "Controls and ingestion paths exist; cross-service recovery and representative outcomes remain open."),
        ("UI-10 Community", "Preserve recipe interactions and gamification while keeping engagement events separate from research labels.", "Existing product workflow; verify separation when reporting results."),
        ("UI-11 Thesis evidence", "Map all 14 Lead chapters to artifact IDs, source routes, and honest readiness states. A planned capture is not a completed reviewed artifact.", "Workspace mapping exists; source-backed captures remain open."),
    ]
    for ident, requirement, status in requirements:
        h(d, ident, 2)
        p(d, requirement)
        p(d, "Current state: " + status)
    h(d, "4 Non-functional requirements")
    table(d, ["ID", "Acceptance condition"], [
        ("NQ-01 Safety", "No name heuristic, missing allergen flag, model confidence, or compound overlap may upgrade a result to Safe. Missing or invalid policy yields Check; conflicts yield Blocked."),
        ("NQ-02 Evidence", "Displayed metrics and explanations retain source, status, protocol, and limitations. A displayed hash is metadata until its artifact has been independently checked."),
        ("NQ-03 Degraded operation", "Missing AI or data services produce explicit pending or error states. Ordinary production requests never silently switch to demo data."),
        ("NQ-04 Privacy and access", "Profile data, uploads, and protected evaluation use application authentication boundaries. Private images and profile values are not published as open evidence."),
        ("NQ-05 Accessibility", "Provide labels, keyboard operation, visible focus, and text cues in addition to color. Large-graph, mobile, and assistive-technology evidence remains required."),
        ("NQ-06 Bounded work", "Graph depth 0–2 and limit 1–500; image cap 10 MiB; model requests support cancellation and a 30-second upstream timeout. Verify on deployment."),
    ])
    h(d, "5 External interfaces and data contracts")
    table(d, ["Interface", "Current contract"], [
        ("Graph", "GET /knowledge/graph using q or root, depth, and limit. Response includes bounded nodes, edges, and coverage. Seed edges are substitution only; confidence may be null."),
        ("Profile", "PUT /auth/settings/cooking and PUT /auth/update persist normalized allergen flags under the authenticated user."),
        ("Substitution", "POST /api/v1/suggest_substitutions carries ingredient, reason, recipe_context, dietary_tags, allergen_flags, and session_id. Candidate safety, evidence, and receipt fields depend on the source."),
        ("Compounds", "GET /api/v1/compound/profile/{ingredient} and POST /api/v1/compound/analyze-pair require source and grounding before factual display."),
        ("Photo", "POST /api/ingredient-detection and POST /api/photo-intelligence/ingredient-recipes use bounded same-origin adapters and caller authentication. A 503 indicates unconfigured integration."),
        ("Evaluation", "The frontend /admin/evaluation route is protected and validates the supplied manifest structure. Schema acceptance does not establish scientific validity."),
    ])
    h(d, "6 Backlog traceability and acceptance gates")
    table(d, ["Gate", "Implemented surface and remaining evidence"], [
        ("G1 and G6 Evaluation", "Dashboard, schema checks, provenance captions, and hidden pending values exist. Close only with verified Lead exports, artifact hashes, strata, intervals, abstention, and review."),
        ("G3 Compound", "Missing values stay unavailable and source identity is retained. Close with official index payload captures and an ethics-reviewed usefulness study."),
        ("G4 Safety", "Profile forwarding and three-valued UI exist. Close with reviewed policy, deployed end-to-end flow, and adjudicated benchmark."),
        ("G5 Graph", "Explorer and bounded API work on seed data. Close with rights, schema, hash, staged import, rollback, coverage, and API-backed captures for the official graph."),
        ("G7 Feedback", "Controls, receipts, and ingestion exist. Close with cross-service choices and outcomes, retries, replay, recovery, and representative evidence."),
        ("G8 Photo", "Camera UI and adapters exist. Close with accepted M3/M4 model or index, provider credentials, rights-cleared images, and real-device evidence."),
        ("G9 Thesis", "Fourteen-chapter workspace exists. Close with source-backed figures, deployment and cost records, ethics/recruitment, and review."),
        ("G10 Voice", "Optional scaffold is deferred and is not a Sprint 2 acceptance item."),
    ])
    h(d, "7 Progress claim rule")
    p(d, "Lead v109 verifies secondary Stage 3 through exact step 327/561 and records only a Segment 6 launch toward step 374. It does not verify step 374 completion, the final adapter, inference, scoring, or model quality. Likewise, a browser graph image proves the seeded UI rather than the canonical research graph, and a scan screenshot proves controls rather than live user-photo recognition.")
    d.save(OUT / "VoTrungTin_IRON_CHEF_Sprint2_SRS_EN.docx")


def case(doc, ident, title, actors, precondition, trigger, steps, alternatives, outcome, acceptance, mapping):
    h(doc, f"{ident} {title}")
    table(doc, ["Field", "Description"], [
        ("Actors", actors), ("Precondition", precondition), ("Trigger", trigger),
        ("Outcome", outcome), ("Traceability", mapping),
    ], keep_together=True)
    h(doc, "Main flow", 2)
    numbered(doc, steps)
    h(doc, "Alternatives and failures", 2)
    bullets(doc, alternatives)
    p(doc, "Acceptance: " + acceptance)


def build_cases():
    d = base("Sprint 2 User Use Cases", "Detailed User Creator and Reviewer journeys with honest acceptance states")
    p(d, "These use cases specify intended behavior at the Member–Lead boundary. A screen or adapter can be implemented while a model, reviewed safety policy, official graph, or research export remains pending. The state of each journey is therefore recorded alongside its expected flow.")
    sources(d)
    h(d, "Use case index")
    table(d, ["Case", "Journey and current boundary"], [
        ("UC-01", "Creator authors a recipe; saved ingredient context still needs cross-service verification."),
        ("UC-02", "User scans or uploads ingredients; editing detections and accepted models remain open."),
        ("UC-03", "User requests a safe substitution; reviewed production policy remains open."),
        ("UC-04", "User explores the graph; UI works with seed data, official import remains open."),
        ("UC-05", "User records cooking feedback; full outcome and recovery evidence remains open."),
        ("UC-06", "Reviewer inspects evaluation; dashboard works, verified Lead exports remain open."),
        ("UC-07", "Reviewer tracks thesis evidence; 14-chapter mapping exists, captures remain open."),
    ])
    case(d, "UC-01", "Creator authors a recipe", "Authenticated Creator", "The Creator may create or edit a recipe.", "Select Create or Edit.", [
        "Enter title, ingredients, quantities, cooking steps, and media.",
        "Save a draft; the system validates fields and preserves the ingredient list.",
        "Preview, revise, and publish the recipe.",
        "Provide stored recipe ingredients as context when a User asks for a substitution."], [
        "On invalid fields or failed save, show the error and retain the draft input.",
        "On insufficient permissions, reject the edit or publication action."],
        "A saved or published recipe retains ingredients and quantities.",
        "Reload preserves ingredient identities and amounts; the substitution request receives the correct recipe context. Community engagement is not recorded as a research label.", "UI-01, UI-10; Member product baseline")
    case(d, "UC-02", "User scans and reviews ingredients", "User, browser, configured detector", "The /scan route opens; image submission has a valid caller identity.", "Choose camera or upload.", [
        "Request camera permission only after an explicit action, or open image selection.",
        "Validate non-empty image type and size before same-origin submission.",
        "If an accepted detector responds, show name, location, confidence, and actual source.",
        "Let the User correct or remove detections before recipe matching or substitution.",
        "Continue only with a reviewed ingredient list."], [
        "Permission denied or camera unsupported: show the reason and offer upload or manual entry.",
        "Unconfigured detector: return 503 integration pending; never invent live detections.",
        "Invalid, empty, or >10 MiB image: return 400/413 without calling a model.",
        "Open UI gap: the observed scan page has no per-detection edit control; that acceptance step still needs implementation and verification."],
        "The reviewed list proceeds, or the journey stops with a clear state.",
        "Browser evidence covers controls only. Verify permission refusal, camera switching, stream cleanup, mobile devices, and a Lead-accepted detector.", "UI-02; G8; Sprint 2 sections 2.2–2.4")
    case(d, "UC-03", "User requests a safety-aware substitution", "User, AI service, reviewed policy provider", "An ingredient and context exist; an empty profile never defaults to Safe.", "Select an ingredient and reason for substitution.", [
        "Collect ingredient, reason, recipe context, dietary tags, saved allergen flags, and session identity.",
        "Request candidates and preserve their receipt and evidence fields when returned.",
        "Resolve local profile conflicts and source policy into Safe, Check, or Blocked.",
        "Display candidate identity, quantity ratio, confidence, source, and supported compound or nutrition evidence. Show missing measurements as unavailable.",
        "Allow selection of a non-Blocked candidate or decline all candidates; record real receipts when available."], [
        "Missing, malformed, or uncertain policy yields Check; absence of an allergen name does not imply Safe.",
        "A conflict yields Blocked and disables the primary action.",
        "Service failure is visible and never replaced by unlabeled demo output.",
        "Missing compound or nutrition measurements are not displayed as measured zero."],
        "The User chooses or declines with safety state and limitations visible.",
        "Test authenticated profile → policy → card end to end, including custom and empty profiles, correct sources and receipts, and a reviewed Lead policy.", "UI-03–06, UI-09; G3/G4/G7")
    case(d, "UC-04", "User explores ingredient relationships", "Guest or User, Knowledge API", "The graph API returns a seed or an accepted imported collection and reports coverage.", "Open /explore/graph or search an ingredient.", [
        "Load a bounded graph response and show its node, edge, and coverage counts.",
        "Search names or aliases and center the selected result.",
        "Drag nodes, pan or zoom, and select nodes or edges for details.",
        "Show relationship type, context, ratio, allergen flags, and auxiliary evidence only when supplied."], [
        "No links: show zero rather than generate edges.",
        "Missing FooDB or USDA data: display unavailable or pending; missing allergen flags do not imply Safe.",
        "The seed API emits substitution only, even though other filter labels appear in the UI.",
        "API failure: show an actionable error and retry state."],
        "The User understands the returned relationships and partial coverage.",
        "The observed seed has 25 nodes and four edges; searching Ginger centers an isolated result. The official Lead graph import remains gate G5.", "UI-07; Epic 1/10; G5")
    case(d, "UC-05", "User records cooking feedback", "User, frontend, monolith, AI ingestion", "A cooking session and valid candidate receipt exist where the event is candidate-specific.", "Accept, reject, skip, or finish a cooking session.", [
        "Create a choice or post-session outcome event.",
        "Send session ID, candidate receipt, and event identity to the backend.",
        "Persist the event durably and project it through the agreed ingestion path.",
        "Show success or an actionable retry state."], [
        "Never fabricate a missing receipt.",
        "Retry or replay must not count an event twice, and a failed projection must not be reported as successful learning.",
        "Likes, XP, and streaks are product interactions, not held-out evaluation labels."],
        "The UI accurately reports whether the feedback was stored.",
        "Exercise accept/reject/skip/taste across services, retries, idempotency, failure and recovery. Do not claim model improvement without representative outcomes.", "UI-09–10; G7")
    case(d, "UC-06", "Reviewer inspects sourced evaluation", "Authorized Reviewer or Admin", "The reviewer is authenticated for /admin/evaluation.", "Open the dashboard or export a figure.", [
        "Load bundled or configured remote Lead manifest data.",
        "Validate structure, states, metrics, metadata, and internal consistency.",
        "Show complete, pending, and rejected records honestly; hide pending numeric values and preserve negative findings.",
        "Export charts with the provenance and limitations actually supplied."], [
        "Unauthenticated visitor is denied or redirected to sign-in.",
        "Malformed manifest or failed remote source fails closed without old scores.",
        "Different datasets or protocols are not presented as a matched winner comparison; a displayed hash does not verify an artifact."],
        "The reviewer sees a truthful dashboard or export.",
        "Verify actual Lead artifacts, dataset split, protocol, seeds, strata, intervals, abstention, and review before treating a score as complete.", "UI-08; G1/G6")
    case(d, "UC-07", "Reviewer tracks thesis evidence", "Reviewer, Member, Lead", "The /thesis workspace maps the 14 Lead chapters.", "Open a chapter evidence item.", [
        "Open the chapter and unique artifact ID.",
        "Read source route and pending, ready-to-capture, captured, or reviewed state.",
        "Capture the real source when available; record caption, date, environment, and version.",
        "Mark reviewed only after examining the evidence and claim limits."], [
        "Missing or unaccepted source remains pending.",
        "Seed/demo screenshots or planned deployment figures do not become production evidence."],
        "Each item has an honest state and traceable source.",
        "Every submitted figure has source, date, version, access boundary, and reviewer. No chapter is marked finished from a plan alone.", "UI-11; G2/G9")
    d.save(OUT / "VoTrungTin_IRON_CHEF_Sprint2_Use_Cases_EN.docx")


def build_data_flow():
    d = base("Sprint 2 Data Flow Specification", "Context diagram, level one flows, data contracts, and trust boundaries")
    p(d, "This document traces data among the browser, frontend adapters, monolith, and Lead-owned AI or evidence sources. Each flow identifies inputs, transformations, stores, and failure states. A diagram arrow expresses an interface or requirement, not proof that an external model is accepted.")
    sources(d)
    h(d, "1 Notation and trust boundaries")
    table(d, ["Term", "Meaning"], [
        ("External entity", "User, Creator, Reviewer, or Lead source that supplies or receives data."),
        ("Process", "UI or service step that validates and transforms input."),
        ("Store", "Profile, recipe, graph, receipt/event, or evidence manifest data."),
        ("Trust boundary", "Browser → same-origin frontend → monolith or model provider. User bearer credentials are not forwarded directly to a model provider."),
    ])
    h(d, "2 Context diagram")
    ctx = draw_flow("dfd-context-en", "IRON CHEF context flow", [
        ("user",90,170,"User Creator\nBrowser UI","entity"),("app",680,170,"ChefKix\nFE and monolith","process"),("ai",1280,170,"Lead AI\nservices","entity"),
        ("review",90,405,"Reviewer\nDashboard","entity"),("store",680,405,"Profile Recipe\nGraph Feedback","store"),("lead",1280,405,"Lead evidence\nmanifest model policy","store")],
        [("user","app","request / event"),("ai","app","model response"),("review","store","review"),("lead","store","provenance")])
    d.add_picture(str(ctx), width=Cm(13.5))
    d.paragraphs[-1].paragraph_format.keep_with_next = True
    p(d, "Figure 1. Logical context; detailed request and response paths are specified below.")
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.paragraphs[-1].runs[0].font.size = Pt(9)
    d.paragraphs[-1].runs[0].italic = True
    h(d, "3 Level one processes")
    table(d, ["Process", "Inputs and outputs"], [
        ("P1 Recipe and profile", "Creator recipe and User allergen settings → stored recipe, normalized profile, and context for P3."),
        ("P2 Photo detection", "Camera/upload image plus caller identity → sourced detections or explicit integration-pending state."),
        ("P3 Substitution and safety", "Ingredient, recipe context, tags, flags, candidate policy → three-valued card and real candidate receipt."),
        ("P4 Graph and compounds", "Search/root/depth/limit or ingredient pair → bounded graph with coverage and grounded evidence where available."),
        ("P5 Feedback", "Session, candidate receipt, choice/outcome → durable event, projection, ingestion, and retry identity."),
        ("P6 Evaluation and thesis", "Lead manifest and artifact references → protected dashboard, figure caption, and evidence readiness state."),
    ])
    h(d, "4 DF 01 Photo to ingredients and recipe candidates")
    photo = draw_flow("dfd-photo-en", "DF 01 Photo to recipe candidates", [
        ("browser",70,170,"Browser\nimage / consent","entity"),("adapter",670,170,"FE adapter\nauth and validate","process"),("model",1270,170,"Lead detector\nif accepted","entity"),
        ("review",1270,405,"User review\nedit pending","process"),("match",670,405,"Recipe match\nadapter","process"),("result",70,405,"Recipe cards\nsource and caveat","store")],
        [("browser","adapter","multipart image"),("adapter","model","bounded image"),("model","review","detections"),("review","match","confirmed list"),("match","result","candidates")])
    d.add_picture(str(photo), width=Cm(13.5))
    d.paragraphs[-1].paragraph_format.keep_with_next = True
    p(d, "Figure 2. User correction is an open UI gate and must occur before downstream action in the accepted flow.")
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.paragraphs[-1].runs[0].font.size = Pt(9)
    d.paragraphs[-1].runs[0].italic = True
    table(d, ["Step", "Input, validation, and failure behavior"], [
        ("Browser", "Request camera only after action or accept upload. Denied or unsupported camera offers upload or manual input; a guest screenshot does not prove a device-camera run."),
        ("Frontend adapter", "Authenticate caller; validate multipart image type, non-zero bytes, and ≤10 MiB. Propagate cancellation and cap upstream work at 30 seconds. Return 400/413 or 503 as appropriate."),
        ("Lead detector", "Return name, box, confidence, and source only from an accepted configured model. M3 is not yet accepted; malformed or failed response is an error, never fabricated output."),
        ("User review", "Correct or remove detections and confirm the list. The observed /scan screen does not yet establish this editing control."),
        ("Recipe matching", "Submit no more than 100 ingredient names of up to 120 characters each through the protected bounded route. Accepted M4 model/index and rights evidence remain open."),
    ])
    h(d, "5 DF 02 Profile to substitution and safety")
    p(d, "Authenticated settings → monolith profile store → normalized allergen flags → substitution request with ingredient, reason, recipe_context, dietary_tags, flags, and session_id → candidate/policy/receipt response → frontend resolver → Safe, Check, or Blocked card → selection or decline. Missing or unreviewed policy never becomes Safe. Custom or unknown terms remain Check without adjudication.")
    table(d, ["Data", "Source and handling"], [
        ("Allergen flags", "Saved under the authenticated profile, normalized, forwarded in later requests, and excluded from public evidence logs."),
        ("Policy result", "Reviewed policy source only. Missing or malformed → Check; conflict → Blocked; Safe only with sufficient reviewed evidence."),
        ("Candidate receipt", "Preserve source identity for feedback; never synthesize a missing receipt."),
        ("Compounds and nutrition", "Separate official evidence path with source, grounding, and units. Null means unavailable and never upgrades allergen safety."),
    ])
    h(d, "6 DF 03 Knowledge graph and compounds")
    p(d, "Browser search or root → GET /knowledge/graph with depth 0–2 and limit 1–500 → seed or imported monolith collection → bounded nodes, edges, and coverage → client focus, pan, zoom, and detail. The captured Ginger result has no edge. The four seed edges are substitutions only. Filter labels for chemical similarity and co-occurrence do not create scientific data. Compound profile and pair requests use a separate official source path; ratio is not confidence and overlap does not prove culinary suitability.")
    h(d, "7 DF 04 Feedback and research evidence")
    p(d, "Cooking session → choice/outcome event with candidate receipt → monolith durable write → projection and AI ingestion/replay. Failed writes, retries, and replay retain event identity and cannot be reported as model learning. Separately, Lead dataset/split/protocol/prediction artifacts → manifest and review decision → frontend schema validation → dashboard/export caption → thesis artifact. The frontend validates structure, not the independent identity, rights, review, or correctness of upstream science.")
    table(d, ["Store", "Owner and current state"], [
        ("Profiles and recipes", "Monolith product storage and frontend flows exist; deployed authenticated acceptance remains to be checked."),
        ("Knowledge graph", "Monolith currently serves a 25-ingredient, four-edge seed. Official G5 import is pending."),
        ("Receipts and feedback events", "AI and monolith paths exist; representative cross-service outcome and recovery evidence remains pending."),
        ("Evaluation manifests", "Lead-owned. Frontend schema and captions exist; verified artifacts and scored results remain pending."),
        ("Thesis captures", "Member and Lead reviewers. Fourteen chapters are mapped; individual source-backed captures and reviews remain pending."),
    ])
    h(d, "8 Data safety and closure matrix")
    table(d, ["Flow", "Evidence required before closure"], [
        ("Photo", "Real devices, provider authentication, accepted detector and retrieval model/index manifests, and image rights. Screen controls are not inference evidence."),
        ("Safety", "Reviewed policy, authenticated deployed profile-to-policy-to-UI test, and adjudicated benchmark. Compound overlap and absence of a match do not establish Safe."),
        ("Graph", "Rights/schema/hash-validated import, coverage checks, rollback, and API-backed captures. A seed API is not the canonical graph."),
        ("Evaluation", "Actual Lead artifacts and checked hashes, split/protocol/seeds, intervals, abstention, and review. A training checkpoint is not a quality score."),
    ])
    p(d, "Lead v109 verifies secondary Stage 3 through exact step 327/561 and records a Segment 6 launch toward step 374. Completion of Stage 3, inference, scoring, and model quality remain open.")
    d.save(OUT / "VoTrungTin_IRON_CHEF_Sprint2_Data_Flow_EN.docx")


if __name__ == "__main__":
    build_srs()
    build_cases()
    build_data_flow()
    print(OUT)
