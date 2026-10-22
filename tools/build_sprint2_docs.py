"""Build the Sprint 2 documentation package from the audited October 2026 baseline."""

from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "sprints" / "2026-10-11"
EVIDENCE = ROOT / "output" / "sprints" / "evidence"
OUT.mkdir(parents=True, exist_ok=True)
NAVY = RGBColor(21, 43, 65)
BLUE = RGBColor(18, 82, 121)
GRAY = RGBColor(82, 93, 103)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def base(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.0)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.space_after = Pt(4)
    for name, size, color in [("Title", 21, RGBColor(0, 0, 0)), ("Heading 1", 14, NAVY), ("Heading 2", 11, BLUE)]:
        s = styles[name]
        s.font.name = "Arial"
        s.font.size = Pt(size)
        s.font.bold = name != "Title"
        s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(9 if name != "Title" else 0)
        s.paragraph_format.space_after = Pt(5)
    # Word's built-in Title style can inherit a blue bottom rule from the theme.
    title_ppr = styles["Title"].element.get_or_add_pPr()
    for old in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(old)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    border.append(bottom)
    title_ppr.append(border)
    if "Caption" not in styles:
        styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    styles["Caption"].font.name = "Arial"
    styles["Caption"].font.size = Pt(8)
    styles["Caption"].font.color.rgb = GRAY
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph(subtitle)
    p.style = "Subtitle"
    p = doc.add_paragraph("IRON CHEF v3  |  Võ Trung Tín  |  Cập nhật 11/10/2026")
    p.style = "Caption"
    p.paragraph_format.space_after = Pt(12)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("IRON CHEF • Sprint 2 • 11/10/2026   |   Trang ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def p(doc, text, style=None):
    return doc.add_paragraph(text, style=style)


def h(doc, text, level=1):
    doc.add_heading(text, level)


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
        set_cell_shading(t.rows[0].cells[i], "DCE9F0")
        if widths:
            t.rows[0].cells[i].width = Cm(widths[i])
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        trPr.append(no_split)
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)
    header_pr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)
    return t


def figure(doc, filename, caption):
    path = EVIDENCE / filename
    if not path.exists():
        return
    doc.add_picture(str(path), width=Cm(11.7))
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


SOURCE_ROWS = [
    ("S1", "Sprint 2 report, §2 và bảng phân công", "Phạm vi Tín: SRS UI, Use Case, Data Flow, nghiệm thu scaffold Member; thời gian 23/09–10/10/2026."),
    ("S2", "BACKLOG_MEMBER.md trên nhánh codex/member-sprint12-deliverables", "Trạng thái Epic 1–12 và các cổng G1–G10; mốc biên soạn 06/10/2026."),
    ("S3", "BACKLOG_LEAD.md tại origin/master 8bb8167, v109", "Mốc 09/10/2026: Stage 3 xác minh đến bước 327; Segment 6 mới khởi chạy tới bước 374; chưa có kết quả chất lượng được xác minh."),
    ("S4", "Ảnh browser thật trong output/sprints/evidence", "Chụp 06/10/2026 từ giao diện FE: /explore/graph và /scan; đồ thị lấy API từ seed cục bộ."),
    ("S5", "Mã FE, monolith và tài liệu tích hợp", "Kiểm tra route, API adapter, trạng thái lỗi, giới hạn dữ liệu và các trường đang thiếu."),
]


def add_sources(doc):
    h(doc, "Nguồn và phạm vi xác minh")
    table(doc, ["Mã", "Nguồn", "Dùng để xác định"], SOURCE_ROWS)
    p(doc, "Phân biệt: ảnh browser chứng minh màn hình đã hiển thị ở lần chạy đó; mã nguồn chứng minh luồng được triển khai; backlog ghi nhận tiến độ nhóm. Không nguồn nào riêng lẻ chứng minh triển khai production, chất lượng mô hình, an toàn y tế hoặc kết quả nghiên cứu.")


def build_srs():
    doc = base("Đặc tả yêu cầu phần mềm Sprint 2", "SRS cho giao diện Member và ranh giới tích hợp với công việc Lead")
    p(doc, "Tài liệu này là bản SRS cập nhật cho phần Võ Trung Tín phụ trách trong Sprint 1–2. Màn graph và scan đã có UI chạy trong browser, nhưng tập đồ thị hiện là dữ liệu seed 25 nguyên liệu, 4 quan hệ; mô hình ảnh và đồ thị nghiên cứu chính thức chưa được nghiệm thu. Vì vậy, các yêu cầu được gắn trạng thái và điều kiện chấp nhận riêng, thay vì đồng nhất giao diện đã có với sản phẩm hoàn chỉnh.")
    table(doc, ["Thông tin", "Giá trị"], [
        ("Phiên bản", "1.1 • kiểm tra ngày 11/10/2026"),
        ("Đối tượng", "Tín, Lead, nhóm FE/BE/AI, người hướng dẫn và người kiểm thử"),
        ("Phạm vi", "UI User, Creator, reviewer; graph, scan, substitution, allergen, feedback, evaluation, thesis evidence"),
        ("Ngoài phạm vi nghiệm thu Sprint 2", "Huấn luyện/đánh giá mô hình, dữ liệu canonical, review lâm sàng, triển khai production, nghiên cứu người dùng"),
    ])
    add_sources(doc)
    h(doc, "1 Bối cảnh và người dùng")
    p(doc, "ChefKix là ứng dụng công thức và nấu ăn có các luồng công khai, tài khoản User/Creator và vùng reviewer/admin. Lead phụ trách bằng chứng khoa học, dữ liệu, mô hình và giao thức đánh giá. Member phụ trách bề mặt người dùng, adapter và ghi nhận tình trạng tích hợp. Các vai trò không thể hoán đổi quyền; endpoint cần xác thực phải từ chối khách chưa đăng nhập.")
    table(doc, ["Vai trò", "Mục tiêu", "Quyền và ranh giới"], [
        ("Khách", "Xem welcome, graph và trang scan công khai", "Có thể thấy UI; thao tác cần tài khoản hoặc dữ liệu AI có thể dừng ở trạng thái đăng nhập/chờ tích hợp."),
        ("User", "Lưu dị ứng, xem công thức, scan, chọn thay thế, nấu và phản hồi", "Dữ liệu hồ sơ và hành vi gắn tài khoản; quyết định an toàn chỉ theo chính sách đã review."),
        ("Creator", "Soạn, sửa, đăng công thức và nguyên liệu", "Giữ ngữ cảnh công thức; không biến chỉ số cộng đồng thành nhãn khoa học."),
        ("Reviewer/Admin", "Xem dashboard đánh giá và thesis evidence", "Chỉ xem dữ liệu có provenance, trạng thái, giới hạn; route đánh giá được bảo vệ."),
        ("Lead/AI provider", "Cung cấp dataset, model, manifest và policy", "Chủ sở hữu nghiệm thu khoa học; UI không tự suy diễn kết quả còn thiếu."),
    ])
    h(doc, "2 Phạm vi UI đã quan sát trong browser")
    figure(doc, "sprint2-03-graph-redesigned-top.jpg", "Ảnh 1. /explore/graph: ô tìm kiếm, bộ lọc quan hệ, số lượng coverage và panel chi tiết, chụp 06/10/2026.")
    p(doc, "Màn graph đọc API Knowledge ở cấu hình đã chạy. Tại mốc chụp: 25 nguyên liệu, 7 nguyên liệu có liên kết, 4 quan hệ substitution. Đây là tập seed phục vụ trình diễn luồng; không đại diện cho toàn bộ đồ thị hoặc hóa học thực phẩm của Lead.")
    figure(doc, "sprint2-03-graph-search-centered.jpg", "Ảnh 2. Tìm Ginger đưa nút về trung tâm; không có cạnh hiển thị. Panel ghi rõ dữ liệu FooDB/USDA chưa sẵn.")
    figure(doc, "sprint2-02b-scan-controls.jpg", "Ảnh 3. /scan: camera, tải ảnh và tìm công thức từ ảnh trong browser; trạng thái khách chưa chứng minh detector thực.")
    h(doc, "3 Yêu cầu chức năng")
    reqs = [
        ("UI-01", "Creator recipe", "Creator tạo/sửa/bản nháp/đăng công thức; nguyên liệu và định lượng giữ được trong request thay thế.", "Có workflow nền; nghiệm thu ngữ cảnh liên service còn cần."),
        ("UI-02", "Scan và nhập liệu", "User chụp/tải ảnh; thấy permission, loading, error, pending và demo riêng. Phải duyệt/sửa detection trước downstream; có đường nhập tay.", "Camera/upload/overlay có; màn scan hiện chưa thấy chỉnh sửa từng detection; model thực và thiết bị thật còn mở."),
        ("UI-03", "Thay thế theo ngữ cảnh", "Gửi ingredient, reason, recipe context, dietary tags, allergen flags; hiển thị ứng viên, ratio, confidence, nguồn, trường thiếu riêng.", "Adapter/UI có; nguồn thật và chuỗi receipt cần nghiệm thu."),
        ("UI-04", "Hồ sơ dị ứng", "EU14/FDA9 và cờ tùy biến được lưu qua settings; request sau mang dữ liệu đã lưu.", "UI và persistence có; thử nghiệm end to end authenticated còn mở."),
        ("UI-05", "Ba trạng thái an toàn", "Safe chỉ theo policy đã review; Check nếu thiếu/không chắc/mâu thuẫn; Blocked nếu xung đột. Blocked vô hiệu hóa hành động chính.", "Resolver và test có; policy production chưa nghiệm thu."),
        ("UI-06", "Giải thích hợp chất", "Nêu nguồn, grounding, đơn vị, giá trị thiếu. Overlap chỉ là sự hiện diện hợp chất; không suy ra phù hợp nấu ăn/an toàn.", "Surface có; index FooDB và bằng chứng UI thật chưa đủ."),
        ("UI-07", "Knowledge graph", "Tìm, focus, kéo/pan/zoom, chọn nút/cạnh, tải neighborhood hữu hạn, chỉ vẽ response và nêu coverage.", "UI/API seed hoạt động; graph chính thức chưa nhập."),
        ("UI-08", "Evaluation", "Route admin bảo vệ; pending/rejected/completed tách biệt; giá trị pending ẩn; caption giữ provenance và giới hạn claim.", "Dashboard/schema có; manifest và đánh giá chuẩn còn thiếu."),
        ("UI-09", "Feedback", "Ghi nhận accept/reject/skip/outcome theo session và candidate receipt; lỗi/retry nhìn thấy.", "UI và ingestion có; thử toàn chuỗi, recovery và outcome đại diện còn mở."),
        ("UI-10", "Community", "Like/comment/XP tiếp tục theo sản phẩm, không dùng như nhãn đánh giá khoa học.", "Workflow nền; tách biệt dữ liệu khi xuất báo cáo."),
        ("UI-11", "Thesis evidence", "14 chương có artifact, nguồn, trạng thái ready/pending; chỉ hoàn thành khi đã chụp/duyệt chứng cứ thật.", "Workspace có; nhiều capture còn pending."),
    ]
    table(doc, ["ID", "Nhóm", "Yêu cầu kiểm thử được", "Trạng thái 11/10"], reqs)
    h(doc, "4 Yêu cầu phi chức năng và tiêu chí kiểm thử")
    table(doc, ["ID", "Tiêu chí", "Điều kiện chấp nhận"], [
        ("NQ-01", "An toàn", "Không suy Safe từ tên, confidence, compound overlap hoặc không thấy dị nguyên. Policy thiếu/invalid → Check; conflict → Blocked."),
        ("NQ-02", "Tính đúng của bằng chứng", "Metric phải có trạng thái, nguồn, protocol, giới hạn; hash hiển thị không tự chứng minh đã kiểm chứng artifact."),
        ("NQ-03", "Suy giảm dịch vụ", "Endpoint/model thiếu → integration pending/lỗi có thể hiểu; không âm thầm thay mock. Demo chỉ khi cấu hình rõ và dán nhãn."),
        ("NQ-04", "Riêng tư và phân quyền", "Profile, ảnh và dashboard qua ranh giới auth; không đưa ảnh thô/profile riêng tư vào bằng chứng công khai."),
        ("NQ-05", "Khả năng tiếp cận", "Nhãn, focus, keyboard và trạng thái không chỉ phụ thuộc màu; cần bằng chứng mobile/keyboard trên graph đã nhập."),
        ("NQ-06", "Giới hạn tài nguyên", "Graph depth 0–2, limit 1–500; ảnh tối đa 10 MiB; request model có hủy/timeout 30 giây; kiểm tra trên môi trường triển khai."),
    ])
    h(doc, "5 Giao diện và hợp đồng tích hợp")
    table(doc, ["Luồng", "Giao diện hiện tại", "Điều kiện và dữ liệu trả về"], [
        ("Graph", "GET /knowledge/graph?q=... hoặc root, depth, limit", "Bounded nodes/edges/coverage; seed chỉ có substitution và ratio/context, confidence có thể null."),
        ("Hồ sơ", "PUT /auth/settings/cooking; PUT /auth/update", "Cờ dị ứng normalized và lưu gắn user; verify caller auth."),
        ("Substitution", "POST /api/v1/suggest_substitutions", "Ingredient, reason, context, tags, flags, session; candidate/safety/receipt tùy nguồn."),
        ("Hợp chất", "GET /api/v1/compound/profile/{ingredient}; POST /api/v1/compound/analyze-pair", "Chỉ hiển thị kết quả có source/grounding; thiếu → unavailable."),
        ("Ảnh", "POST /api/ingredient-detection; POST /api/photo-intelligence/ingredient-recipes", "Caller auth, input bound, 503 khi provider chưa cấu hình; không coi adapter là model."),
        ("Đánh giá", "GET /admin/evaluation trong FE", "Protected; parse manifest; chưa chứng minh tính đúng khoa học của nguồn."),
    ])
    h(doc, "6 Ma trận backlog và cổng nghiệm thu")
    table(doc, ["Backlog", "Hiện có", "Cần để đóng"], [
        ("G1/G6 • Evaluation", "Dashboard, schema, caption và ẩn metric pending", "Lead exports thật, artifact hash kiểm chứng, strata/CI/abstention và review."),
        ("G3 • Compound", "UI giữ null/unavailable và source", "Official index/payload chụp thật, nghiên cứu usefulness có đạo đức."),
        ("G4 • Safety", "Profile, forwarding, tri-state", "Policy đã review và e2e deployed; benchmark có adjudication."),
        ("G5 • Graph", "Explorer và bounded API trên seed", "Manifest graph chính thức, rights/hash/schema, staged import, rollback, coverage và ảnh API-backed."),
        ("G7 • Feedback", "Controls, receipts, ingestion", "Accept/reject/skip/taste liên service, retry/replay/recovery, outcome đại diện."),
        ("G8 • Photo", "Camera/UI/adapters", "M3/M4 accepted model/index, provider auth, quyền ảnh và thiết bị thật."),
        ("G9 • Thesis", "Map 14 chương", "Capture nguồn thật, deployment/cost, study/rehearsal."),
        ("G10 • Voice", "Scaffold tùy chọn", "Hoãn; không tính nghiệm thu Sprint 2."),
    ])
    h(doc, "7 Quy tắc báo cáo tiến độ")
    p(doc, "Lead v109 đã xác minh checkpoint Stage 3 đến bước 327/561 và chỉ ghi nhận việc khởi chạy Segment 6 hướng tới bước 374. Không khẳng định bước 374 hoàn thành, adapter cuối, inference, điểm benchmark hay chất lượng model. Ảnh graph chứng minh UI và dữ liệu seed, không chứng minh đồ thị nghiên cứu hoàn chỉnh. Ảnh scan chứng minh controls, không chứng minh detector hoạt động với ảnh người dùng.")
    doc.save(OUT / "VoTrungTin_IRON_CHEF_Sprint2_SRS.docx")


def uc(doc, ident, title, actors, pre, trigger, steps, exceptions, post, acceptance, mapping):
    h(doc, f"{ident} {title}")
    table(doc, ["Mục", "Đặc tả"], [("Tác nhân", actors), ("Tiền điều kiện", pre), ("Kích hoạt", trigger), ("Hậu điều kiện", post), ("Truy vết", mapping)])
    h(doc, "Luồng chính", 2)
    for i, item in enumerate(steps, 1): p(doc, f"{i}. {item}")
    h(doc, "Luồng thay thế và lỗi", 2)
    for item in exceptions: bullet(doc, item)
    p(doc, "Tiêu chí chấp nhận: " + acceptance)


def build_use_cases():
    doc = base("Đặc tả ca sử dụng Sprint 2", "Hành trình User Creator và Reviewer dựa trên UI và backlog hiện tại")
    p(doc, "Bộ ca sử dụng này mô tả hành vi mong muốn và ghi rõ phần có thể xác nhận ở lần chạy browser 06/10/2026. Một luồng có giao diện hoặc adapter chưa đồng nghĩa mô hình AI, chính sách an toàn hay dữ liệu nghiên cứu đã được Lead chấp nhận.")
    add_sources(doc)
    h(doc, "Danh mục ca sử dụng")
    table(doc, ["ID", "Ca sử dụng", "Màn chính", "Tình trạng"], [
        ("UC-01", "Creator quản lý công thức", "/creator và màn recipe", "Workflow nền; liên kết context cần kiểm chứng"),
        ("UC-02", "User chụp/tải ảnh và duyệt nguyên liệu", "/scan", "Controls có; sửa detection/model thật còn mở"),
        ("UC-03", "User yêu cầu thay thế an toàn", "recipe/cook", "UI và tri-state có; policy thật còn mở"),
        ("UC-04", "User khám phá graph và chi tiết nút", "/explore/graph", "Chạy với seed; full import pending"),
        ("UC-05", "User phản hồi sau nấu", "cook", "Controls/receipt có; e2e outcome pending"),
        ("UC-06", "Reviewer xem đánh giá", "/admin/evaluation", "Protected dashboard; metric evidence pending"),
        ("UC-07", "Reviewer quản lý thesis evidence", "/thesis", "Map chương có; capture pending"),
    ])
    uc(doc, "UC-01", "Creator quản lý công thức", "Creator đã đăng nhập", "Có quyền soạn/sửa công thức.", "Chọn tạo hoặc sửa công thức.", [
        "Nhập tiêu đề, thành phần, định lượng, bước nấu và media theo form hiện có.",
        "Lưu bản nháp; hệ thống kiểm tra dữ liệu và giữ nguyên danh sách nguyên liệu.",
        "Xem trước, chỉnh sửa rồi đăng công thức.",
        "Khi User hỏi thay thế, recipe context lấy từ dữ liệu công thức đã lưu."], [
        "Thiếu trường bắt buộc hoặc lỗi lưu: hiển thị lỗi, giữ dữ liệu đang nhập.",
        "Không đủ quyền sửa/xuất bản: chặn và báo trạng thái quyền."],
        "Công thức lưu hoặc được xuất bản, có nguyên liệu và định lượng truy xuất được.",
        "Save/reload giữ ingredient list và quantities; recipe context truyền đúng; tương tác cộng đồng không bị tính là nhãn nghiên cứu.", "UI-01, UI-10; Member baseline/feedback")
    uc(doc, "UC-02", "User chụp tải ảnh và duyệt nguyên liệu", "User; browser; detector service nếu đã cấu hình", "Route /scan mở được; thao tác gửi ảnh cần auth hợp lệ.", "Chọn Use camera hoặc Upload photo.", [
        "Browser xin quyền camera sau thao tác rõ ràng hoặc mở hộp chọn ảnh.",
        "UI kiểm tra loại ảnh, kích thước và gửi qua same-origin adapter.",
        "Nếu có detector hợp lệ, hiển thị tên, vị trí, confidence và source thực tế.",
        "User kiểm tra và sửa tên detection trước khi dùng cho recipe matching hoặc thay thế.",
        "Tiếp tục với danh sách đã xác nhận."], [
        "Từ chối quyền hoặc thiết bị không hỗ trợ: báo lý do, dùng upload/nhập tay.",
        "Không có detector: 503 integration pending; không dựng detection giả.",
        "Ảnh sai định dạng/rỗng/quá 10 MiB: lỗi 400/413; không gọi model.",
        "Điểm hở hiện tại: màn scan đã quan sát chưa có control sửa từng detection; bước xác nhận/sửa và fallback nhập tay cần hoàn thiện/kiểm thử."],
        "Danh sách đã được User xác nhận hoặc luồng dừng rõ ràng; không có output model giả.",
        "Ảnh guest chụp được controls, chưa chứng minh inference; cần kiểm thử permissions, switch, cleanup và thiết bị thật.", "UI-02; G8; Sprint 2 2.2–2.4")
    uc(doc, "UC-03", "User yêu cầu thay thế an toàn", "User; AI service; policy đã review", "Có ingredient và context; profile có thể rỗng nhưng trạng thái không được mặc định Safe.", "Chọn nguyên liệu cần thay và lý do.", [
        "Đọc ingredient, reason, recipe context, dietary tags, allergen flags và session.",
        "Gửi POST /api/v1/suggest_substitutions và nhận candidate/receipt/evidence nếu có.",
        "Áp dụng profile conflict và chính sách nguồn theo Safe, Check hoặc Blocked.",
        "Hiển thị identity, quantity ratio, confidence, source, compound và nutrition với giá trị thiếu là unavailable.",
        "User chọn ứng viên không bị Blocked hoặc từ chối; ghi nhận receipt khi có."], [
        "Policy thiếu/invalid/không chắc: Check; không suy Safe từ vắng tên dị nguyên.",
        "Conflict: Blocked và chặn primary action.",
        "Service lỗi: thông báo rõ; không thay bằng demo không dán nhãn.",
        "Thiếu compound/nutrition: không hiển thị 0 như giá trị đo."],
        "Ứng viên được chọn hoặc từ chối với trạng thái an toàn và giới hạn chứng cứ rõ ràng.",
        "E2E authenticated profile→policy→card, cả custom và empty profile; source và receipt đúng; policy Lead đã review.", "UI-03–06, UI-09; G3/G4/G7")
    uc(doc, "UC-04", "User khám phá đồ thị nguyên liệu", "Khách/User; Knowledge API", "API có thể trả graph seed hoặc collection nhập chính thức; UI nêu coverage.", "Mở /explore/graph hoặc tìm nguyên liệu.", [
        "Tải kết quả hữu hạn từ /knowledge/graph; hiển thị số node/edge/coverage.",
        "Tìm theo tên/alias; nút kết quả được focus vào vùng trung tâm.",
        "Kéo nút, pan/zoom, chọn node hoặc edge để mở panel chi tiết.",
        "Xem loại quan hệ, context, ratio, allergen flag và evidence phụ nếu thực sự có."], [
        "Không có node/link: hiển thị 0 link, không tự tạo cạnh.",
        "FooDB/USDA chưa có: hiển thị unavailable/pending; thiếu cờ dị ứng không có nghĩa Safe.",
        "Filter chemical/co-occurrence có trong UI nhưng API seed chỉ trả substitution; không quảng bá đó là quan hệ đã xác nhận.",
        "API lỗi: panel trạng thái lỗi và retry."],
        "User hiểu quan hệ được trả về và coverage còn thiếu.",
        "Browser seed: 25 node, 4 edge; tìm Ginger focus chính giữa và 0 link. Import Lead chính thức vẫn là cổng G5.", "UI-07; G5; Epic 1/10")
    uc(doc, "UC-05", "User ghi nhận lựa chọn và kết quả nấu", "User; FE; monolith; AI ingestion", "Có cooking session và candidate receipt hợp lệ nếu phản hồi gắn đề xuất.", "Accept/reject/skip hoặc hoàn tất buổi nấu.", [
        "UI tạo event cho lựa chọn hoặc kết quả sau nấu.",
        "Gửi session, candidate receipt và event id qua backend.",
        "Backend ghi bền và chuyển ingestion/projection theo hợp đồng.",
        "UI thông báo thành công hoặc cho retry khi lỗi."], [
        "Thiếu receipt: không dựng receipt; giới hạn feedback theo chính sách ingestion.",
        "Retry/replay: cùng event không tính lặp; failure projection không báo là mô hình đã học.",
        "Like, XP, streak là tín hiệu sản phẩm riêng, không dùng như nhãn benchmark."],
        "Phản hồi có trạng thái lưu/không lưu minh bạch.",
        "Chạy liên service cả accept/reject/skip/taste, lỗi và recovery; chưa có outcome đại diện thì không tuyên bố cải thiện mô hình.", "UI-09–10; G7")
    uc(doc, "UC-06", "Reviewer xem đánh giá có nguồn", "Reviewer/Admin", "Đăng nhập và có quyền truy cập /admin/evaluation.", "Mở dashboard hoặc xuất hình.", [
        "FE tải bundled hoặc configured remote manifest.",
        "Schema validation kiểm tra cấu trúc, trạng thái, metric, metadata và tính nhất quán.",
        "Hiển thị record complete/pending/rejected; ẩn số pending và giữ negative findings.",
        "Xuất biểu đồ cùng caption và provenance được nguồn cung cấp."], [
        "Khách chưa đăng nhập: chuyển sign-in/deny.",
        "Manifest hỏng hoặc remote lỗi: fail closed, không quay lại điểm cũ.",
        "Khác dataset/protocol: không so sánh như matched winner; hash hiển thị chưa chứng minh artifact."],
        "Reviewer có view/export trung thực theo source.",
        "Kiểm chứng Lead artifact, split/protocol/seeds, strata/CI/abstention và review trước khi coi score là hoàn tất.", "UI-08; G1/G6")
    uc(doc, "UC-07", "Reviewer quản lý bằng chứng luận văn", "Reviewer/Tín/Lead", "Có workspace /thesis và mapping 14 chương.", "Xem một mục bằng chứng chương.", [
        "Mở chương và artifact ID duy nhất.",
        "Xem route nguồn và trạng thái pending/ready-to-capture/captured/reviewed.",
        "Chụp từ môi trường thật khi nguồn sẵn và ghi caption, thời điểm, phiên bản.",
        "Đánh dấu reviewed chỉ sau khi kiểm tra bằng chứng và giới hạn claim."], [
        "Nguồn chưa có/không được Lead xác nhận: giữ pending.",
        "Ảnh seed/demo hoặc thông số deployment theo kế hoạch không đổi thành chứng cứ production."],
        "Mục có trạng thái trung thực và truy xuất nguồn.",
        "Từng hình có nguồn, ngày, phiên bản, quyền truy cập và reviewer; 14 chương không có mục bị đánh dấu xong giả.", "UI-11; G2/G9")
    doc.save(OUT / "VoTrungTin_IRON_CHEF_Sprint2_Use_Cases.docx")


def draw_flow(name, title, nodes, edges):
    width, height = 1700, 660
    im = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(im)
    font_path = "C:/Windows/Fonts/arial.ttf"
    bold_path = "C:/Windows/Fonts/arialbd.ttf"
    f = ImageFont.truetype(font_path, 29)
    fb = ImageFont.truetype(bold_path, 34)
    fs = ImageFont.truetype(font_path, 21)
    d.text((50, 30), title, font=fb, fill=(21, 43, 65))
    boxes = {}
    for key, x, y, label, typ in nodes:
        box = (x, y, x + 330, y + 115)
        boxes[key] = box
        fill = (232, 243, 249) if typ == "process" else ((250, 241, 226) if typ == "store" else (238, 243, 238))
        d.rounded_rectangle(box, radius=15, fill=fill, outline=(49, 94, 116), width=3)
        lines = label.split("\n")
        for j, line in enumerate(lines):
            d.text((x + 18, y + 15 + 35 * j), line, font=f if j == 0 else fs, fill=(21, 43, 65))
    for src, dst, label in edges:
        a, b = boxes[src], boxes[dst]
        if abs((a[0]+a[2])/2-(b[0]+b[2])/2) > abs((a[1]+a[3])/2-(b[1]+b[3])/2):
            start = (a[2], (a[1]+a[3])/2) if a[0] < b[0] else (a[0], (a[1]+a[3])/2)
            end = (b[0], (b[1]+b[3])/2) if a[0] < b[0] else (b[2], (b[1]+b[3])/2)
        else:
            start = ((a[0]+a[2])/2, a[3]) if a[1] < b[1] else ((a[0]+a[2])/2, a[1])
            end = ((b[0]+b[2])/2, b[1]) if a[1] < b[1] else ((b[0]+b[2])/2, b[3])
        d.line((start, end), fill=(83, 115, 132), width=5)
        import math
        ang = math.atan2(end[1]-start[1], end[0]-start[0])
        arrow = [(end[0],end[1]),(end[0]-20*math.cos(ang-.5),end[1]-20*math.sin(ang-.5)),(end[0]-20*math.cos(ang+.5),end[1]-20*math.sin(ang+.5))]
        d.polygon(arrow, fill=(83, 115, 132))
        mx, my = (start[0]+end[0])/2, (start[1]+end[1])/2
        bounds = d.textbbox((0, 0), label, font=fs)
        tw = bounds[2] - bounds[0]
        tx, ty = mx - tw / 2, my - 35
        d.rectangle((tx-8,ty-3,tx+tw+8,ty+28),fill="white")
        d.text((tx,ty),label,font=fs,fill=(55,75,90))
    path = OUT / f"{name}.png"
    im.save(path)
    return path


def diagram(doc, path, caption):
    doc.add_picture(str(path), width=Cm(15.4))
    x = doc.add_paragraph(caption, style="Caption")
    x.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_data_flow():
    doc = base("Đặc tả luồng dữ liệu Sprint 2", "Sơ đồ ngữ cảnh DFD và hợp đồng dữ liệu cho các luồng Member Lead")
    p(doc, "Tài liệu mô tả dữ liệu đi qua browser, FE adapter, monolith và AI/Lead; mỗi luồng có nguồn, phép biến đổi, kho lưu và nhánh lỗi. Đường mũi tên là hợp đồng hiện tại hoặc yêu cầu rõ trạng thái, không hàm ý endpoint AI đã có mô hình được chấp nhận.")
    add_sources(doc)
    h(doc, "Ký hiệu và ranh giới tin cậy")
    table(doc, ["Loại", "Ý nghĩa"], [
        ("External entity", "User, Creator, Reviewer hoặc Lead cung cấp/nhận dữ liệu."),
        ("Process", "UI hoặc service xử lý input, kiểm tra và chuyển đổi."),
        ("Store", "Profile, recipe, graph seed/import, event/receipt hoặc evidence manifest."),
        ("Trust boundary", "Browser ↔ same-origin FE ↔ monolith/AI provider; token không truyền thẳng đến model provider."),
    ])
    ctx = draw_flow("dfd-context", "DFD ngữ cảnh IRON CHEF", [
        ("user",90,170,"User Creator\nBrowser UI","entity"),("app",680,170,"ChefKix\nFE và monolith","process"),("ai",1280,170,"AI Lead\nservices","entity"),
        ("review",90,405,"Reviewer\nDashboard","entity"),("store",680,405,"Profile Recipe\nGraph Feedback","store"),("lead",1280,405,"Lead evidence\nmanifest model policy","store")],
        [("user","app","request / event"),("ai","app","model response"),("review","store","review"),("lead","store","provenance")])
    diagram(doc, ctx, "Hình 1. Ngữ cảnh logic. Dòng hai chiều thực tế được diễn giải trong bảng chi tiết bên dưới.")
    doc.add_page_break()
    h(doc, "DFD mức 1 các xử lý chính")
    table(doc, ["Mã", "Xử lý", "Đầu vào", "Đầu ra và kho"], [
        ("P1", "Quản lý công thức/hồ sơ", "Creator recipe; User allergen settings", "Recipe store; profile normalized; context và flags cho P3."),
        ("P2", "Nhận diện ảnh", "Camera/upload image; caller auth", "Detection list với source/confidence hoặc integration-pending; không lưu ảnh public."),
        ("P3", "Thay thế/an toàn", "Ingredient, recipe context, tags, flags, candidate policy", "Safe/Check/Blocked card, candidate receipt; missing fields unavailable."),
        ("P4", "Graph/compound", "Search/root/depth/limit; ingredient pair", "Bounded node/edge, coverage, grounded compound evidence; seed khác full import."),
        ("P5", "Feedback", "Session, candidate receipt, choice/outcome", "Durable event, projection/ingestion; retry và idempotency."),
        ("P6", "Evaluation/thesis", "Lead manifest/export và artifact references", "Protected dashboard, caption/figure, evidence state."),
    ])
    h(doc, "DF 01 Ảnh đến nguyên liệu và công thức")
    pic = draw_flow("dfd-photo", "DF 01 Ảnh đến công thức", [
        ("browser",70,170,"Browser\nimage / consent","entity"),("adapter",670,170,"FE adapter\nauth + validate","process"),("model",1270,170,"Lead detector\nif accepted","entity"),
        ("review",1270,405,"User review\nedit pending","process"),("match",670,405,"Recipe match\nadapter","process"),("result",70,405,"Recipe cards\nsource + caveat","store")],
        [("browser","adapter","multipart image"),("adapter","model","bounded image"),("model","review","detections"),("review","match","confirmed list"),("match","result","candidates")])
    diagram(doc,pic,"Hình 2. Luồng photo. Bước User review/edit còn là gap UI; không được bỏ qua khi mô tả acceptance.")
    table(doc,["Bước", "Dữ liệu và kiểm tra", "Thất bại / giới hạn"],[
        ("1 Browser", "Camera sau action hoặc upload; client giữ ảnh tạm.", "Denied/unsupported → upload hoặc nhập tay; không claim camera thật từ ảnh guest."),
        ("2 FE adapter", "Verify caller qua monolith; multipart image, type/size >0 và ≤10 MiB; hủy/timeout 30s.", "400/413 invalid/oversize; 503 nếu chưa cấu hình; không fallback mock âm thầm."),
        ("3 AI detector", "Trả detections gồm name, bounding box, confidence, source khi model accepted.", "M3 chưa accepted; response hỏng/upstream lỗi phải là lỗi rõ."),
        ("4 Review", "User sửa/loại detection rồi xác nhận input.", "Control edit hiện chưa xác nhận trên /scan; cần triển khai và test."),
        ("5 Match", "Danh sách tên ≤100, mỗi tên ≤120 ký tự; route matching được bảo vệ và bounded.", "M4 rejected; provider/index accepted và rights pending."),
    ])
    h(doc, "DF 02 Hồ sơ dị ứng đến quyết định thay thế")
    p(doc, "User settings → monolith profile store → FE đọc allergen flags → POST suggest_substitutions với ingredient, reason, recipe_context, dietary_tags, allergen_flags và session_id → AI candidates/policy/receipt → FE resolver → Safe/Check/Blocked card → lựa chọn hoặc từ chối. Profile flags không được biến thành chứng nhận Safe nếu không có policy đã review. Cờ custom hoặc unknown phải giữ trạng thái Check khi thiếu adjudication.")
    table(doc,["Dữ liệu", "Nguồn", "Quy tắc hiển thị / lưu"],[
        ("Allergen flags", "Authenticated profile/settings", "Persist normalized; forward trong request; không log vào public evidence."),
        ("Policy result", "Reviewed AI/policy endpoint", "Missing/malformed → Check; conflict → Blocked; Safe chỉ nếu đủ chứng cứ."),
        ("Candidate receipt", "AI response", "Giữ identity nguồn để feedback; không bịa receipt."),
        ("Compounds/nutrition", "Official source endpoints", "Đường riêng, source/grounding/units; null → unavailable, không nâng cấp an toàn."),
    ])
    h(doc, "DF 03 Knowledge graph và hợp chất")
    p(doc, "Browser search/root → GET /knowledge/graph (depth 0–2, limit 1–500) → monolith seed hoặc graph import → nodes/edges/coverage → client focus/pan/zoom/detail. Tìm Ginger trong ảnh browser trả nút không cạnh. Bốn cạnh seed chỉ là substitution; UI filter chemical similarity và co-occurrence không làm các loại quan hệ đó xuất hiện. Profile/pair compound là request riêng sang nguồn chính thức, chỉ hiển thị khi có grounding và source. Ratio substitution không phải confidence; overlap hợp chất không chứng minh thay thế nấu ăn.")
    h(doc, "DF 04 Phản hồi và bằng chứng đánh giá")
    p(doc, "User cooking session → event với choice/outcome/candidate receipt → monolith durable write → message/projection → AI ingestion/replay. Lỗi ghi, retry hoặc replay phải giữ event identity và không báo nhầm là model đã học. Lead dataset/split/protocol/prediction artifacts → manifest/review decision → FE schema validation → dashboard/export caption → thesis artifact. FE có thể xác thực cấu trúc, không thể tự kiểm chứng hash, review, quyền dữ liệu hoặc tính đúng của score.")
    table(doc,["Kho / dữ liệu", "Chủ sở hữu", "Trạng thái hiện tại"],[
        ("Profile / recipe", "Monolith product", "Lưu trữ và FE flow có; e2e deployed cần kiểm chứng."),
        ("Graph", "Monolith + Lead data", "Seed 25 nguyên liệu/4 cạnh; full import G5 pending."),
        ("Candidate receipts/events", "AI + monolith", "Đường code có; G7 cross-service outcome/recovery pending."),
        ("Evaluation manifests", "Lead", "FE schema/caption có; Lead exports được xác minh và scores còn pending."),
        ("Thesis captures", "Member/Lead reviewer", "14 chương mapping; ảnh nguồn thật và review từng artifact còn pending."),
    ])
    h(doc, "Ma trận an toàn dữ liệu và điều kiện đóng")
    table(doc,["Luồng", "Không được suy diễn", "Bằng chứng đóng"],[
        ("Ảnh", "Ảnh controls ≠ detector hoạt động; model candidate ≠ accepted.", "Thiết bị thật, provider auth, model/index manifest và quyền ảnh."),
        ("Safety", "Không thấy dị nguyên ≠ Safe; compound overlap ≠ an toàn.", "Policy reviewed, deployed e2e và adjudicated benchmark."),
        ("Graph", "Seed API ≠ canonical graph; filter UI ≠ quan hệ khoa học.", "Import hash/rights/schema, coverage, rollback và capture API-backed."),
        ("Evaluation", "Training checkpoint ≠ metric/chất lượng; metadata hash ≠ artifact verification.", "Lead actual exports, split/protocol/seeds, review, intervals và source artifacts."),
    ])
    p(doc, "Mốc Lead v109 09/10/2026 chỉ xác minh Stage 3 exact step 327/561 và launch Segment 6 hướng step 374. Điều kiện hoàn tất Stage 3, inference, scoring, model quality và các cổng dữ liệu/safety bên trên chưa đóng.")
    doc.save(OUT / "VoTrungTin_IRON_CHEF_Sprint2_Data_Flow.docx")


if __name__ == "__main__":
    build_srs()
    build_use_cases()
    build_data_flow()
    print(OUT)
