from __future__ import annotations

import importlib.util
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(__file__).resolve().parents[2]
BASE_PATH = ROOT / "tmp" / "source-inspect" / "build_sprint_reports.py"
spec = importlib.util.spec_from_file_location("base_builder", BASE_PATH)
base = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(base)

OUT_DIR = ROOT / "output" / "kltn_sprints"
OUT_DIR.mkdir(parents=True, exist_ok=True)


def report(
    n: int,
    period: str,
    date: str,
    summary: list[tuple[str, str]],
    objective: str,
    scope: str,
    items: list[tuple[str, str, str]],
    features: list[tuple[str, str, str, str]],
    evidence: list[str],
    risk: str,
    mitigation: str,
    next_backend: str,
    next_frontend: str,
    assignments: list[tuple[str, str, str]],
    outline_item: str,
    outline_status: str,
    outline_note: str,
    conclusion: str,
    next_title: str | None = None,
    next_assignment_title: str | None = None,
) -> dict:
    return locals()


REPORTS = [
    report(
        1,
        "07/09/2026 – 22/09/2026",
        "22/09/2026",
        [
            ("Định hướng nghiên cứu:", "Xác định IRON CHEF là hệ thống gợi ý thay thế nguyên liệu dựa trên đồ thị tri thức đa nguồn, tập trung vào khả năng tổng quát hóa trên unseen pairs, an toàn dị ứng và giải thích hóa học thực phẩm."),
            ("Đánh giá nền tảng:", "Rà soát kết quả Đồ án 2, kiến trúc ChefKix hiện có, các module cần mở rộng AI và các nền tảng đã có như AI Service, knowledge graph, feedback telemetry và thesis evidence workspace."),
            ("Tổ chức thực hiện:", "Tách backlog thành hai track Lead và Member theo đúng IRON CHEF v3, ưu tiên các Epic P0/P1 làm nền cho đề tài và đặt Epic 12 Voice-Vision Copilot ở mức tùy chọn."),
        ],
        "Chuyển đề cương khóa luận thành câu hỏi nghiên cứu, backlog có thứ tự ưu tiên, kế hoạch dữ liệu/mô hình và lịch trình có thể theo dõi.",
        "Đánh giá baseline ChefKix, xác định novelty của IRON CHEF, lập backlog Lead/Member, Gantt chart và kế hoạch thu thập bằng chứng cho các chương khóa luận.",
        [
            ("1.1", "Đánh giá kết quả Đồ án 2 và xác định các điểm tích hợp AI cần mở rộng", "Đang thực hiện"),
            ("1.2", "Chốt câu hỏi nghiên cứu về repeated pairs, unseen pairs, safety và explainability", "Đã xác định"),
            ("1.3", "Phân tách Lead Backlog và Member Backlog theo Epic, ưu tiên P0/P1/P2", "Đã lập"),
            ("1.4", "Lập Gantt chart và kế hoạch thu thập bằng chứng cho khóa luận", "Đang hoàn thiện"),
        ],
        [
            ("Research", "Research questions và novelty", "Đã xác định", "Tập trung vào unseen-pair generalization, safety và chemistry-grounded explanation"),
            ("Planning", "Lead/Member backlog", "Đã lập", "Epic 1–9 và Epic 1–12 được phân vai song song"),
            ("Platform", "ChefKix baseline", "Đang rà soát", "Xác định module, API và giao diện cần tích hợp AI"),
        ],
        [
            "Đề cương chi tiết KLTN v5 và lịch trình từ 07/09/2026 đến 17/01/2027.",
            "BACKLOG_LEAD.md và BACKLOG_MEMBER.md với các Epic, priority và deliverable liên quan.",
            "Sơ đồ repository ChefKix và thesis evidence workspace dùng để theo dõi bằng chứng theo chương.",
        ],
        "Phạm vi khóa luận bao gồm đồ thị, mô hình, safety, explainability, computer vision, frontend và triển khai nên dễ vượt quá thời gian.",
        "Bám Priority Guide của backlog: bảo vệ P0 HGAT/benchmark/chemical engine, sau đó P1 safety/behavioral; chỉ làm P2 và Epic 12 khi các đầu ra thesis-critical đã ổn định.",
        "Khảo sát MISKG, Recipe1MSubs, FooDB, FlavorDB, USDA, Open Food Facts và FDA; viết SRS, Use Case và Data Flow.",
        "Chuẩn bị graph explorer, evaluation dashboard, allergen profile và camera scaffold theo Member Epic 1–4.",
        [
            ("Phan Phú Thọ", "Lead track: graph, HGAT, benchmark, chemical reasoning, safety và kế hoạch dữ liệu", "ML and Data Science"),
            ("Võ Trung Tín", "Member track: frontend intelligence surfaces, camera, dashboard, allergen profile và evidence UI", "Systems and Frontend Intelligence"),
        ],
        "Đánh giá và Lên kế hoạch (07/09 – 22/09)",
        "Đang thực hiện",
        "Đã có định hướng nghiên cứu và backlog; Gantt chart cùng baseline evidence tiếp tục được hoàn thiện.",
        "Sprint 1 đặt lại toàn bộ phạm vi theo khóa luận IRON CHEF thay vì báo cáo phát triển ChefKix chung. Các sprint sau sẽ theo đề cương, backlog và đầu ra nghiên cứu đã chốt.",
    ),
    report(
        2,
        "23/09/2026 – 10/10/2026",
        "10/10/2026",
        [
            ("Dữ liệu nghiên cứu:", "Khảo sát MISKG cho quan hệ thay thế, Recipe1MSubs cho bối cảnh công thức, FooDB và FlavorDB cho hợp chất, USDA FoodData Central cho dinh dưỡng, cùng Open Food Facts và dữ liệu cảnh báo FDA cho safety."),
            ("Yêu cầu hệ thống:", "Đặc tả các luồng nhận diện nguyên liệu từ ảnh/văn bản, gợi ý thay thế theo ngữ cảnh, cảnh báo dị ứng, giải thích hợp chất, gamification và tương tác cộng đồng trên ChefKix."),
            ("Đánh giá:", "Đưa repeated pairs và unseen pairs thành hai lát cắt bắt buộc; xác định các metric Hit@1, Hit@5, Hit@10, MRR, NDCG và tỷ lệ vi phạm dị ứng cần thu thập."),
        ],
        "Hoàn thành tài liệu đặc tả yêu cầu và ma trận dữ liệu để mọi Epic trong backlog có đầu vào, tiêu chí chấp nhận và nguồn bằng chứng rõ ràng.",
        "Khảo sát dataset, SRS, Use Case, Data Flow, yêu cầu chức năng/phi chức năng và hợp đồng dữ liệu cho substitution, allergen profile và compound explanation.",
        [
            ("2.1", "Lập ma trận dataset, provenance, schema và phạm vi sử dụng", "Theo kế hoạch"),
            ("2.2", "Đặc tả SRS cho Creator, User và các luồng nghiên cứu", "Theo kế hoạch"),
            ("2.3", "Xây dựng Use Case và Data Flow cho AI pipeline và ChefKix", "Theo kế hoạch"),
            ("2.4", "Định nghĩa API contract và acceptance criteria cho safety/explanation", "Theo kế hoạch"),
        ],
        [
            ("Data", "MISKG, Recipe1MSubs, FooDB, FlavorDB, USDA, OFF, FDA", "Theo kế hoạch", "Mỗi nguồn có vai trò, provenance và mapping rõ ràng"),
            ("Product", "Creator và User journeys", "Theo kế hoạch", "Bao phủ đăng công thức, scan, substitution, safety và chemistry"),
            ("Evaluation", "Repeated vs unseen protocol", "Theo kế hoạch", "Metric và baseline được chốt trước khi train"),
        ],
        [
            "Dataset matrix và tài liệu SRS của khóa luận.",
            "Use Case, Data Flow và backlog acceptance criteria cho Lead/Member.",
            "API contract sơ bộ cho substitution response, allergen flags, nutritional delta và compound explanation.",
        ],
        "Tên nguyên liệu và schema giữa MISKG, Recipe1MSubs, FooDB, USDA và Open Food Facts không đồng nhất; dữ liệu có thể có thiếu hoặc trùng.",
        "Xây dựng canonical ingredient vocabulary, giữ provenance theo từng nguồn, kiểm tra license và tách rõ dữ liệu dùng cho train, benchmark và safety.",
        "Thiết kế schema đồ thị dị thể, kiến trúc HGAT/Feature-Adaptive Gating, AI Service contract và integration boundary với Modular Monolith.",
        "Dựng UI/UX trên Figma cho scan, substitution card, safety indicators; khởi tạo graph explorer, evaluation dashboard, allergen profile và camera scaffold.",
        [
            ("Phan Phú Thọ", "Dataset matrix, canonical vocabulary, evaluation protocol và AI API requirements", "Lead requirements"),
            ("Võ Trung Tín", "Use Case, Data Flow, SRS UI requirements và Member scaffold acceptance", "Member requirements"),
        ],
        "Phân tích và Đặc tả yêu cầu (23/09 – 10/10)",
        "Theo kế hoạch",
        "Đầu ra cần đạt là SRS, backlog cập nhật, Use Case và Data Flow đủ chi tiết để chuyển sang thiết kế hệ thống.",
        "Sprint 2 làm rõ IRON CHEF cần chứng minh điều gì và dữ liệu nào đủ để chứng minh. Đây là điểm kiểm soát để tránh xây giao diện hoặc mô hình không phục vụ câu hỏi nghiên cứu.",
    ),
    report(
        3,
        "11/10/2026 – 28/10/2026",
        "28/10/2026",
        [
            ("Đồ thị và mô hình:", "Thiết kế schema HeteroData với substitution edges, compound edges và recipe co-occurrence edges; xác định ba nhóm đặc trưng Epicure/semantic, USDA/nutritional và FooDB/chemical."),
            ("AI Service:", "Thiết kế ranh giới FastAPI, model registry, ONNX Runtime và các endpoint cho prediction, copilot, detection, moderation, safety và compound explanation; backend ChefKix gọi qua contract có version."),
            ("Giao diện nghiên cứu:", "Thiết kế graph explorer, benchmark dashboard, compound explanation card, allergen safety UI và camera flow; mọi surface có trạng thái pending-data theo thesis evidence manifest."),
        ],
        "Chốt thiết kế đủ chi tiết để Lead có thể xây graph/model và Member có thể dựng các bề mặt hiển thị trí tuệ AI song song.",
        "Graph schema, HGAT architecture, Feature-Adaptive Gating, FastAPI API contract, Modular Monolith integration, Figma UI/UX và evidence manifest.",
        [
            ("3.1", "Thiết kế schema đồ thị đa nguồn và kiến trúc HGAT", "Theo kế hoạch"),
            ("3.2", "Thiết kế AI Service contract và integration boundary", "Theo kế hoạch"),
            ("3.3", "Thiết kế UI/UX cho graph, dashboard, safety, chemistry và camera", "Theo kế hoạch"),
            ("3.4", "Tạo fixtures/mock data để Member phát triển không bị block bởi model", "Theo kế hoạch"),
        ],
        [
            ("Model", "Heterogeneous graph + adaptive gating", "Theo kế hoạch", "Đặc trưng và relation types được xác định trước train"),
            ("Service", "FastAPI và model registry", "Theo kế hoạch", "Response version có safety, explanation và provenance"),
            ("Frontend", "Graph, evaluation, camera và safety surfaces", "Theo kế hoạch", "Dùng mock data trước, thay bằng export thật sau"),
        ],
        [
            "Bản thiết kế đồ thị, model architecture và API contract.",
            "Figma prototype và component inventory cho các Member Epic 1–6.",
            "Thesis evidence manifest cho Chapters 5, 6, 7, 8, 10 và 11.",
        ],
        "Lead và Member có thể hiểu khác nhau về payload hoặc trạng thái chưa có dữ liệu, dẫn đến phải sửa UI khi model hoàn tất.",
        "Version hóa contract, dùng mock fixtures có schema giống response thật và biểu diễn rõ ready, pending-data và capture-needed.",
        "Bắt đầu tải dữ liệu, canonical vocabulary, graph reconstruction và baseline HGAT; xuất graph sample, compound index và allergen index.",
        "Hoàn thiện graph renderer, node/edge detail, search/filter, dashboard shell, allergen profile và camera permission/bounding-box flow.",
        [
            ("Phan Phú Thọ", "Graph schema, HGAT architecture, API contract và model registry design", "Lead architecture"),
            ("Võ Trung Tín", "Figma, graph explorer, evaluation dashboard, safety UI và camera scaffolds", "Member architecture"),
        ],
        "Thiết kế Hệ thống và UI/UX (11/10 – 28/10)",
        "Theo kế hoạch",
        "Đầu ra gồm bản thiết kế đồ thị/mô hình, kiến trúc hệ thống/API và UI/UX hoàn chỉnh theo đề cương.",
        "Sprint 3 là mốc khóa thiết kế: từ đây, code và thí nghiệm sẽ được đánh giá theo cùng một graph schema, API contract và evidence plan.",
    ),
    report(
        4,
        "29/10/2026 – 12/11/2026",
        "12/11/2026",
        [
            ("Thu thập và chuẩn hóa:", "Tải và kiểm tra MISKG, Recipe1MSubs, FooDB, FlavorDB, Epicure-Core, USDA, Open Food Facts và dữ liệu FDA; xây canonical vocabulary nối tên nguyên liệu giữa các nguồn."),
            ("Graph reconstruction:", "Xây dựng HeteroData với substitution, compound và co-occurrence relations; chuẩn bị graph sample 500 ingredients và các index cho Member phát triển UI."),
            ("Baseline:", "Đối chiếu graph/HGAT v1 đã có trong backlog với mục tiêu v2; giữ checkpoint và model registry làm baseline để không nhầm giữa code đã phục vụ và kết quả thesis mới."),
        ],
        "Tạo được nền dữ liệu và graph reconstruction có provenance, làm đầu vào cho HGAT v2 và các tầng giải thích/an toàn.",
        "Lead Epic 1–2: data acquisition, canonical vocabulary, multi-source graph, triple-encoded features và các export mà Member sử dụng.",
        [
            ("4.1", "Tải, kiểm tra và ghi nhận provenance các dataset", "Theo kế hoạch"),
            ("4.2", "Map canonical ingredient vocabulary giữa các nguồn", "Theo kế hoạch"),
            ("4.3", "Xây HeteroData và relation types cho graph v2", "Theo kế hoạch"),
            ("4.4", "Xuất graph_sample, compound_index và allergen_index", "Theo kế hoạch"),
        ],
        [
            ("Data", "Dataset acquisition and provenance", "Theo kế hoạch", "Giữ nguồn và phiên bản cho khả năng tái lập"),
            ("Graph", "Multi-source HeteroData", "Theo kế hoạch", "Substitution, chemical và recipe relations"),
            ("Baseline", "HGAT v1 versus v2", "Đã có nền tảng", "HGAT v1 đã serving; v2 là phần nghiên cứu cần đánh giá"),
        ],
        [
            "Raw/normalized data manifest và canonical vocabulary report.",
            "HeteroData graph statistics, sample graph và export JSON cho frontend.",
            "Model registry và checkpoint baseline HGAT v1 để so sánh với v2.",
        ],
        "Dữ liệu lớn, file FooDB/Open Food Facts và graph construction có thể vượt RAM hoặc thời gian xử lý cục bộ.",
        "Chunked processing, OOM-safe pipeline, lưu intermediate artifacts, checksum dữ liệu và chỉ dùng graph sample cho UI trong giai đoạn đầu.",
        "Hoàn thiện triple-encoded features, stratified split, HGAT v2 training plan và Symbolic Safety Layer/Selective Abstention.",
        "Dùng graph_sample, compound index và allergen index để hoàn thiện node/edge panels, search/filter và mock data replacement.",
        [
            ("Phan Phú Thọ", "Dataset pipeline, vocabulary, HeteroData graph và HGAT baseline comparison", "Lead graph reconstruction"),
            ("Võ Trung Tín", "Graph explorer integration, detail panels, search/filter và data fixtures", "Member graph surface"),
        ],
        "Cài đặt và Phát triển Phase 1 - Dữ liệu và Mô hình AI (29/10 – 26/11)",
        "Đang triển khai",
        "Sprint 4 hoàn thành phần đầu của Phase 1; graph v2 và các export được chuẩn bị để tiếp tục trong Sprint 5.",
        "Sprint 4 phân biệt rõ baseline HGAT v1 đã có với đóng góp mới của khóa luận: graph đa nguồn, feature gating, unseen-pair benchmark, safety và explainability.",
    ),
    report(
        5,
        "13/11/2026 – 26/11/2026",
        "26/11/2026",
        [
            ("Multi-signal features:", "Ghép Epicure embeddings, vector dinh dưỡng USDA và fingerprint hợp chất FooDB thành feature groups; bổ sung relation-aware input cho graph v2."),
            ("Safety layer:", "Thiết kế allergen constraint checker theo FDA Big 9 và EU 14, nutritional validation, hard filter và Selective Abstention khi confidence hoặc safety evidence không đạt."),
            ("Exports cho Member:", "Đóng gói ingredient-to-compound và ingredient-to-allergen indices, response schema và sample data để UI chemistry/safety có thể tích hợp trước khi benchmark hoàn tất."),
        ],
        "Hoàn tất Phase 1 ở mức có thể train/evaluate: graph và feature contract ổn định, safety layer có quy tắc xác định và output cho frontend có provenance.",
        "Lead Epic 2 và Epic 5: multi-source graph, feature-adaptive input, allergen index, nutritional validation, allergen-safe filter và selective abstention.",
        [
            ("5.1", "Hoàn thiện triple-encoded node features và graph v2 input", "Theo kế hoạch"),
            ("5.2", "Xây allergen constraint checker và nutritional validation", "Theo kế hoạch"),
            ("5.3", "Cài đặt hard safety filter và Selective Abstention", "Theo kế hoạch"),
            ("5.4", "Đóng gói compound/allergen exports cho frontend", "Theo kế hoạch"),
        ],
        [
            ("Model input", "Semantic, nutritional, chemical signals", "Theo kế hoạch", "Sẵn sàng cho HGAT v2 và ablation"),
            ("Safety", "FDA Big 9, EU 14 và hard filter", "Theo kế hoạch", "Không đưa gợi ý vi phạm thành kết quả chính"),
            ("Explainability", "Compound/allergen export", "Theo kế hoạch", "Có tên compound, concentration và flags khi nguồn có dữ liệu"),
        ],
        [
            "Feature schema, graph statistics và data quality report.",
            "Allergen index, safety rule documentation và test cases cho blocked substitutions.",
            "Compound index/API fixture để Member Epic 5–6 dùng trước khi có response production.",
        ],
        "Safety layer có thể đánh dấu thiếu dữ liệu là an toàn nếu quy tắc thiếu trạng thái unknown; feature groups cũng có thể bị lệch scale.",
        "Tách rõ safe, check, blocked và unknown; ưu tiên fail closed; chuẩn hóa feature, ghi lại provenance và bổ sung unit tests cho edge cases.",
        "Train HGAT v2 trên MISKG split, chạy baselines, metric Hit@K/MRR/NDCG, ablation và chemical compound reasoning engine.",
        "Wire compound explanation card, safety indicators, blocked reason và comparison view vào CookingPlayer/recipe detail.",
        [
            ("Phan Phú Thọ", "Feature pipeline, Symbolic Safety Layer, Selective Abstention và export contracts", "Lead safety and data"),
            ("Võ Trung Tín", "Compound explanation UI, safety states, blocked reason và comparison view", "Member trust surfaces"),
        ],
        "Cài đặt và Phát triển Phase 1 - Dữ liệu và Mô hình AI (29/10 – 26/11)",
        "Theo kế hoạch",
        "Đầu ra cần có graph hoàn chỉnh, checkpoint đầu tiên, module an toàn dị ứng và tài liệu để bước sang benchmark/tích hợp.",
        "Sprint 5 biến graph thành đầu vào có cấu trúc cho nghiên cứu và biến safety/explainability thành contract mà người dùng có thể nhìn thấy, không chỉ là logic ẩn trong backend.",
    ),
    report(
        6,
        "27/11/2026 – 10/12/2026",
        "10/12/2026",
        [
            ("HGAT v2 và benchmark:", "Huấn luyện trên MISKG theo split stratified, đo Hit@1/5/10, MRR và NDCG, so sánh PMI, Epicure cosine, GISMo, Gemini zero-shot và Mistral theo protocol đã chốt."),
            ("Ablation:", "Tách chemical-only, nutritional-only, semantic-only và all-signals để xác định nguồn tri thức nào đóng góp cho unseen-pair generalization và tránh negative transfer."),
            ("Chemical reasoning:", "Xây fingerprint/index từ FooDB và FlavorDB, tính shared-compound percentage, nutritional delta và sinh giải thích có bằng chứng cho substitution pair."),
        ],
        "Tạo số liệu nghiên cứu có thể tái lập và kiểm tra đóng góp của multi-signal HGAT, đồng thời hoàn thiện tầng giải thích hóa thực phẩm.",
        "Lead Epic 3–4: training, benchmark, baselines, ablation, compound fingerprint, shared-compound calculator và compound-explainer model registry.",
        [
            ("6.1", "Train HGAT v2 và đánh giá trên repeated/unseen pairs", "Theo kế hoạch"),
            ("6.2", "Chạy baseline và ghi Hit@K, MRR, NDCG", "Theo kế hoạch"),
            ("6.3", "Chạy ablation theo semantic/nutritional/chemical/all signals", "Theo kế hoạch"),
            ("6.4", "Xây compound reasoning engine và response explanation", "Theo kế hoạch"),
        ],
        [
            ("Benchmark", "HGAT v2 versus baselines", "Theo kế hoạch", "Có benchmark_results và bảng repeated/unseen"),
            ("Ablation", "Signal contribution", "Theo kế hoạch", "Có ablation_results cho dashboard"),
            ("Chemistry", "Shared compounds và nutritional delta", "Theo kế hoạch", "Có explanation payload có provenance"),
        ],
        [
            "Benchmark results, ablation results và model/config manifest.",
            "Compound fingerprint index và ví dụ giải thích substitution pair.",
            "Checkpoint/model registry record cho chefkix:hgat-v2 và compound-explainer-v1.",
        ],
        "Kết quả benchmark có thể thấp hơn baseline; GPU budget và chất lượng mapping compound có thể làm thí nghiệm không ổn định.",
        "Đóng băng dataset/split/config, lưu seed và logs, chạy baseline trước, dùng early stopping/checkpoint và chỉ kết luận theo kết quả đo được.",
        "Đóng gói AI Service bằng FastAPI/ONNX Runtime, tích hợp YOLOv8 + CLIP, allergen guard, nutrition validation và substitution API.",
        "Load benchmark/ablation JSON vào evaluation dashboard, hiển thị chemistry explanation và chuẩn bị graph/recipe flow cho dữ liệu thật.",
        [
            ("Phan Phú Thọ", "HGAT v2 training, benchmark protocol, ablation và compound reasoning", "Thesis centerpiece"),
            ("Võ Trung Tín", "Evaluation dashboard, benchmark table, ablation charts và compound explanation UI", "Thesis evidence UI"),
        ],
        "Cài đặt và Phát triển Phase 2 - AI Service và Tích hợp (27/11 – 24/12)",
        "Đang triển khai",
        "Sprint 6 tạo phần thực nghiệm cốt lõi của Phase 2; các kết quả cần được đóng gói để tích hợp vào AI Service và dashboard.",
        "Sprint 6 là trung tâm khoa học của khóa luận: mọi tuyên bố về cải thiện phải được gắn với split, baseline, ablation và bảng số liệu cụ thể.",
    ),
    report(
        7,
        "11/12/2026 – 24/12/2026",
        "24/12/2026",
        [
            ("AI Service:", "Đóng gói pipeline FastAPI, model registry, ONNX Runtime, YOLOv8/CLIP và substitution endpoints; giữ timeout, rate limit, correlation id và graceful degradation."),
            ("Safety integration:", "Đưa allergen guard thành filter bắt buộc trong substitution flow; trả về Safe, Check, Blocked hoặc Abstain kèm nguyên nhân và evidence thay vì để LLM quyết định tự do."),
            ("Product integration:", "Nối photo intelligence, graph explorer, evaluation dashboard, compound explanation, allergen profile và feedback instrument vào ChefKix Modular Monolith/Next.js."),
        ],
        "Hoàn thiện pipeline từ nhận diện ảnh/văn bản tới gợi ý, kiểm định an toàn, giải thích và hiển thị trên ChefKix.",
        "Lead Epic 5–8 và Member Epic 5–10: AI Service, allergen guard, detection, behavioral feedback, flavor pairing, photo pipeline, dashboard và graph real-data wiring.",
        [
            ("7.1", "Đóng gói FastAPI, ONNX Runtime, model registry và versioned endpoints", "Theo kế hoạch"),
            ("7.2", "Tích hợp allergen guard, selective abstention và compound explanation", "Theo kế hoạch"),
            ("7.3", "Wire YOLOv8/CLIP photo pipeline và graph explorer dữ liệu thật", "Theo kế hoạch"),
            ("7.4", "Kết nối feedback instrument và evaluation dashboard", "Theo kế hoạch"),
        ],
        [
            ("AI Service", "FastAPI + ONNX + model registry", "Theo kế hoạch", "Có endpoint và fallback theo contract"),
            ("Safety", "Safe/Check/Blocked/Abstain", "Theo kế hoạch", "Safety guard là constraint bắt buộc"),
            ("Product", "Photo, graph, dashboard và feedback", "Theo kế hoạch", "Tích hợp các surface thesis-critical"),
        ],
        [
            "API smoke tests, model registry manifest và integration logs.",
            "Ảnh chụp scan → recipe → substitution → safety → explanation flow.",
            "Evaluation dashboard, graph explorer và thesis evidence capture briefs.",
        ],
        "Các mô hình/dịch vụ có độ trễ và lỗi khác nhau; nếu UI hiển thị fallback như kết quả đã được xác minh sẽ làm sai thông điệp khoa học.",
        "Tách pending-data và verified, hiển thị source/version, timeout/fallback rõ ràng, fail closed cho safety và giữ correlation id xuyên suốt pipeline.",
        "Chạy unit/integration test, đánh giá repeated/unseen và allergen benchmark, tối ưu API/page load và chuẩn bị deployment.",
        "Chạy visual QA toàn bộ sản phẩm, hoàn thiện capture cho Chapters 5–8/10–11 và chuẩn bị user study.",
        [
            ("Phan Phú Thọ", "AI Service packaging, allergen guard, detection endpoints, feedback và integration tests", "AI integration"),
            ("Võ Trung Tín", "Photo pipeline, safety/chemistry UI, dashboard, graph explorer và evidence capture", "Product integration"),
        ],
        "Cài đặt và Phát triển Phase 2 - AI Service và Tích hợp (27/11 – 24/12)",
        "Theo kế hoạch",
        "Đầu ra kỳ vọng là AI Service ổn định và ChefKix tích hợp các tính năng AI theo đúng đề cương.",
        "Sprint 7 nối phần nghiên cứu với sản phẩm. Các màn hình không chỉ để demo mà phải giữ được provenance, safety state và giới hạn bằng chứng theo thesis manifest.",
    ),
    report(
        8,
        "25/12/2026 – 17/01/2027",
        "17/01/2027",
        [
            ("Kiểm thử và đánh giá:", "Chạy unit/integration test cho software, đánh giá model trên repeated/unseen pairs, benchmark allergen safety, ablation và User Study về tính hữu ích của giải thích hóa học."),
            ("Hiệu năng và triển khai:", "Đo thời gian pipeline từ nhận diện tới gợi ý an toàn, tối ưu page/API load, đóng gói Docker và triển khai server/cloud theo mục tiêu phản hồi dưới 3 giây khi điều kiện tài nguyên cho phép."),
            ("Hồ sơ khóa luận:", "Hoàn thiện các chương 5, 6, 7, 8, 10, 11; xuất hình dashboard/graph, viết báo cáo tổng kết, slide, video demo và tập dượt bảo vệ."),
        ],
        "Đóng gói bằng chứng nghiên cứu và sản phẩm để hệ thống có thể nghiệm thu, triển khai trình diễn và bảo vệ khóa luận.",
        "Kiểm thử, experimental report, performance, deployment, thesis chapters, presentation, video demo và handover; Epic 11 là bắt buộc, Epic 12 chỉ làm nếu còn thời gian.",
        [
            ("8.1", "Unit/integration test và regression cho các luồng phần mềm", "Theo kế hoạch"),
            ("8.2", "Đánh giá Hit@K/MRR/NDCG, ablation, safety benchmark và User Study", "Theo kế hoạch"),
            ("8.3", "Tối ưu API/page load, Docker và deployment readiness", "Theo kế hoạch"),
            ("8.4", "Hoàn thiện thesis evidence, report, slides, video và rehearsal", "Theo kế hoạch"),
        ],
        [
            ("Evaluation", "Repeated/unseen, baselines và ablation", "Theo kế hoạch", "Kết luận chỉ dựa trên số liệu và protocol đã đóng băng"),
            ("Safety", "Allergen violation rate và selective abstention", "Theo kế hoạch", "Có bảng head-to-head với GPT-4o/Gemini theo thiết kế nghiên cứu"),
            ("Delivery", "Deployment và thesis defense", "Theo kế hoạch", "Có báo cáo, slide, video và checklist bàn giao"),
        ],
        [
            "Test reports, benchmark tables, allergen safety tables và User Study summary.",
            "Ảnh thesis-ready từ compound explanation, allergen safety, feedback, photo pipeline, graph và evaluation dashboard.",
            "Bản báo cáo KLTN, slide thuyết trình và video demo theo lịch trình đề cương.",
        ],
        "Thời gian cuối kỳ ngắn trong khi phải vừa chạy thực nghiệm vừa viết báo cáo; kết quả có thể chưa đủ để chứng minh mọi mục tiêu.",
        "Đóng băng scope P0/P1, ghi rõ giới hạn, dùng evidence manifest để ưu tiên figure/bảng bắt buộc, tạo backup demo và không dùng Epic 12 để đánh đổi thesis evidence.",
        "Sau khi nộp và bảo vệ: sửa lỗi theo phản hồi, bổ sung documentation và lập kế hoạch mở rộng nutrition mapping, cascade substitution và user feedback learning.",
        "Nếu P0/P1 đã ổn định, cân nhắc Epic 12 Voice-Vision Copilot như hướng phát triển dài hạn; không đưa vào kết luận nếu chưa có bằng chứng.",
        [
            ("Phan Phú Thọ", "Thực nghiệm, benchmark, safety evaluation, deployment và các chương nghiên cứu", "Lead thesis and release"),
            ("Võ Trung Tín", "User Study UI, visual/evidence capture, slides, video demo và rehearsal", "Member thesis and demo"),
        ],
        "Kiểm thử & Đánh giá (25/12 – 08/01); Triển khai & Báo cáo (09/01 – 17/01)",
        "Theo kế hoạch",
        "Sprint 8 bao phủ hai chặng cuối của đề cương: kiểm thử/đánh giá và triển khai/báo cáo; hoàn thiện và bảo vệ diễn ra sau đó.",
        "Tám sprint được khép lại bằng bằng chứng, không chỉ bằng tính năng. IRON CHEF cần bảo vệ được cả câu hỏi nghiên cứu, safety claim, explainability claim và khả năng tích hợp sản phẩm trên ChefKix.",
        next_title="7. KẾ HOẠCH SAU SPRINT 8",
        next_assignment_title="8. PHÂN CÔNG GIAI ĐOẠN HOÀN THIỆN VÀ BẢO VỆ",
    ),
]


def fill_docx(report_data: dict, out_path: Path):
    doc = Document(str(base.REFERENCE_DOCX))
    p = doc.paragraphs
    base.set_heading(p[0], f"BÁO CÁO TIẾN ĐỘ KHÓA LUẬN TỐT NGHIỆP (SPRINT {report_data['n']})", 2)
    title = "IRON CHEF - Hệ thống gợi ý thay thế nguyên liệu ẩm thực dựa trên đồ thị tri thức đa nguồn, kiểm soát an toàn dị ứng và giải thích hóa học thực phẩm"
    base.fill_paragraph(p[1], [("Đề tài:", True, False), (f" {title}", False, False)])
    base.fill_paragraph(p[2], [("Trường:", True, False), (" Trường Đại học Công nghệ Thông tin - ĐHQG TP.HCM", False, False)])
    base.fill_paragraph(p[3], [("Giảng viên hướng dẫn:", True, False), (" ThS. Trần Thị Hồng Yến", False, False)])
    base.fill_paragraph(p[4], [("Sinh viên thực hiện:", True, False), (" Phan Phú Thọ - 23521520, Võ Trung Tín - 23521595", False, False)])
    base.fill_paragraph(p[5], [("Thời gian sprint:", True, False), (f" {report_data['period']} | Báo cáo: {report_data['date']}", False, False)])
    base.set_heading(p[6], "1. TÓM TẮT ĐIỀU HÀNH", 3)
    for idx, (label, text) in enumerate(report_data["summary"], start=7):
        base.fill_paragraph(p[idx], [(label, True, False), (f" {text}", False, False)])
    base.set_heading(p[10], "2. MỤC TIÊU & PHẠM VI", 3)
    base.fill_paragraph(p[11], [("Mục tiêu:", True, False), (f" {report_data['objective']}", False, False)])
    base.fill_paragraph(p[12], [("Phạm vi:", True, False), (f" {report_data['scope']}", False, False)])
    base.set_heading(p[13], "3. CÁC HẠNG MỤC TRỌNG TÂM", 3)
    base.set_heading(p[14], "4. BẢNG TRẠNG THÁI BACKLOG", 3)
    base.set_heading(p[15], "5. ĐẦU RA VÀ MINH CHỨNG", 3)
    for idx, text in enumerate(report_data["evidence"], start=16):
        base.fill_paragraph(p[idx], [(text, False, False)])
    base.set_heading(p[19], "6. RỦI RO & BIỆN PHÁP", 3)
    base.set_heading(p[20], report_data.get("next_title") or f"7. KẾ HOẠCH 2 TUẦN TỚI (SPRINT {report_data['n'] + 1})", 3)
    base.set_heading(p[21], "7.1 Nghiên cứu, Backend và AI", 4)
    base.fill_paragraph(p[22], [(report_data["next_backend"], False, False)])
    base.set_heading(p[23], "7.2 Frontend, QA và Evidence", 4)
    base.fill_paragraph(p[24], [(report_data["next_frontend"], False, False)])
    base.set_heading(p[25], report_data.get("next_assignment_title") or f"8. PHÂN CÔNG (SPRINT {report_data['n'] + 1})", 3)
    base.set_heading(p[26], "9. ĐỐI CHIẾU VỚI ĐỀ CƯƠNG", 3)
    base.set_heading(p[27], "10. KẾT LUẬN", 3)
    base.fill_paragraph(p[28], [(report_data["conclusion"], False, False)])
    base.set_heading(p[29], "11. Phụ lục A - Liên kết và Minh hoạ", 3)
    base.fill_paragraph(p[30], [("Nguồn nội dung:", True, False), (" Đề cương chi tiết KLTN v5, BACKLOG_LEAD.md, BACKLOG_MEMBER.md và thesis evidence workspace.", False, False)])
    base.fill_paragraph(p[31], [("Người lập báo cáo:", True, False), (" Phan Phú Thọ, Võ Trung Tín", False, False)])

    tables = doc.tables
    headers = [
        ["Mục", "Hạng mục trọng tâm", "Trạng thái"],
        ["Nhánh", "Backlog / Epic", "Trạng thái", "Ghi chú"],
        ["Rủi ro", "Biện pháp khắc phục"],
        ["Thành viên", "Phụ trách", "Ghi chú"],
        ["Hạng mục trong Đề cương", "Trạng thái thực tế", "Ghi chú/Mức độ hoàn thành"],
    ]
    for ti, row in enumerate(headers):
        for ci, value in enumerate(row):
            base.set_cell(tables[ti].cell(0, ci), value, bold=True)
    for ri, row in enumerate(report_data["items"], start=1):
        for ci, value in enumerate(row):
            base.set_cell(tables[0].cell(ri, ci), value)
    for ri, row in enumerate(report_data["features"], start=1):
        for ci, value in enumerate(row):
            base.set_cell(tables[1].cell(ri, ci), value)
    base.set_cell(tables[2].cell(1, 0), report_data["risk"])
    base.set_cell(tables[2].cell(1, 1), report_data["mitigation"])
    for ri, row in enumerate(report_data["assignments"], start=1):
        for ci, value in enumerate(row):
            base.set_cell(tables[3].cell(ri, ci), value, bold=(ci == 0))
    for ci, value in enumerate([report_data["outline_item"], report_data["outline_status"], report_data["outline_note"]]):
        base.set_cell(tables[4].cell(1, ci), value, bold=(ci == 1))
    doc.save(str(out_path))


def build_pdf(report_data: dict, out_path: Path):
    styles = base.pdf_styles()
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import inch
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer

    doc = BaseDocTemplate(
        str(out_path), pagesize=A4,
        leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch,
        title=f"IRON CHEF - Sprint {report_data['n']}", author="Phan Phú Thọ, Võ Trung Tín",
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame])])
    esc = base.esc
    story = []
    story.append(Paragraph(f"BÁO CÁO TIẾN ĐỘ KHÓA LUẬN TỐT NGHIỆP (SPRINT {report_data['n']})", styles["title"]))
    title = "IRON CHEF - Hệ thống gợi ý thay thế nguyên liệu ẩm thực dựa trên đồ thị tri thức đa nguồn, kiểm soát an toàn dị ứng và giải thích hóa học thực phẩm"
    metadata = [
        ("Đề tài:", title),
        ("Trường:", "Trường Đại học Công nghệ Thông tin - ĐHQG TP.HCM"),
        ("Giảng viên hướng dẫn:", "ThS. Trần Thị Hồng Yến"),
        ("Sinh viên thực hiện:", "Phan Phú Thọ - 23521520, Võ Trung Tín - 23521595"),
        ("Thời gian sprint:", f"{report_data['period']} | Báo cáo: {report_data['date']}"),
    ]
    for label, value in metadata:
        story.append(Paragraph(f"<b>{esc(label)}</b> {esc(value)}", styles["body"]))
    story.append(Paragraph("1. TÓM TẮT ĐIỀU HÀNH", styles["h3"]))
    for label, value in report_data["summary"]:
        story.append(Paragraph(f"<b>{esc(label)}</b> {esc(value)}", styles["body"]))
    story.append(Paragraph("2. MỤC TIÊU &amp; PHẠM VI", styles["h3"]))
    story.append(Paragraph(f"• <b>Mục tiêu:</b> {esc(report_data['objective'])}", styles["body"]))
    story.append(Paragraph(f"• <b>Phạm vi:</b> {esc(report_data['scope'])}", styles["body"]))
    story.append(Paragraph("3. CÁC HẠNG MỤC TRỌNG TÂM", styles["h3"]))
    story.append(base.table_flow([["Mục", "Hạng mục trọng tâm", "Trạng thái"], *report_data["items"]], styles, [0.55 * inch, 3.35 * inch, 1.6 * inch]))
    story.append(Spacer(1, 10))
    story.append(Paragraph("4. BẢNG TRẠNG THÁI BACKLOG", styles["h3"]))
    story.append(base.table_flow([["Nhánh", "Backlog / Epic", "Trạng thái", "Ghi chú"], *report_data["features"]], styles, [1.0 * inch, 1.8 * inch, 1.0 * inch, 2.2 * inch]))
    story.append(Spacer(1, 10))
    story.append(Paragraph("5. ĐẦU RA VÀ MINH CHỨNG", styles["h3"]))
    for value in report_data["evidence"]:
        story.append(Paragraph(f"• {esc(value)}", styles["body"]))
    story.append(Paragraph("6. RỦI RO &amp; BIỆN PHÁP", styles["h3"]))
    story.append(base.table_flow([["Rủi ro", "Biện pháp khắc phục"], [report_data["risk"], report_data["mitigation"]]], styles, [3.25 * inch, 3.25 * inch]))
    story.append(Spacer(1, 10))
    story.append(Paragraph(esc(report_data.get("next_title") or f"7. KẾ HOẠCH 2 TUẦN TỚI (SPRINT {report_data['n'] + 1})"), styles["h3"]))
    story.append(Paragraph("7.1 Nghiên cứu, Backend và AI", styles["h4"]))
    story.append(Paragraph(esc(report_data["next_backend"]), styles["body"]))
    story.append(Paragraph("7.2 Frontend, QA và Evidence", styles["h4"]))
    story.append(Paragraph(esc(report_data["next_frontend"]), styles["body"]))
    story.append(Paragraph(esc(report_data.get("next_assignment_title") or f"8. PHÂN CÔNG (SPRINT {report_data['n'] + 1})"), styles["h3"]))
    story.append(base.table_flow([["Thành viên", "Phụ trách", "Ghi chú"], *report_data["assignments"]], styles, [1.7 * inch, 3.35 * inch, 1.45 * inch], bold_cells={(1, 0), (2, 0)}))
    story.append(Spacer(1, 10))
    story.append(Paragraph("9. ĐỐI CHIẾU VỚI ĐỀ CƯƠNG", styles["h3"]))
    story.append(base.table_flow([["Hạng mục trong Đề cương", "Trạng thái thực tế", "Ghi chú/Mức độ hoàn thành"], [report_data["outline_item"], report_data["outline_status"], report_data["outline_note"]]], styles, [2.35 * inch, 1.55 * inch, 2.6 * inch], bold_cells={(1, 1)}))
    story.append(Spacer(1, 10))
    story.append(Paragraph("10. KẾT LUẬN", styles["h3"]))
    story.append(Paragraph(esc(report_data["conclusion"]), styles["body"]))
    story.append(Paragraph("11. Phụ lục A - Liên kết và Minh hoạ", styles["h3"]))
    story.append(Paragraph("<b>Nguồn nội dung:</b> Đề cương chi tiết KLTN v5, BACKLOG_LEAD.md, BACKLOG_MEMBER.md và thesis evidence workspace.", styles["body"]))
    story.append(Paragraph("<b>Người lập báo cáo:</b> Phan Phú Thọ, Võ Trung Tín", styles["body"]))
    doc.build(story)


def main():
    base.extract_fonts()
    base.pdf_font_setup()
    manifest = []
    for data in REPORTS:
        stem = f"PhanPhuTho_VoTrungTin_IRON_CHEF_Sprint{data['n']}"
        docx = OUT_DIR / f"{stem}.docx"
        pdf = OUT_DIR / f"{stem}.pdf"
        fill_docx(data, docx)
        build_pdf(data, pdf)
        manifest.append({"sprint": data["n"], "docx": str(docx), "pdf": str(pdf)})
    (OUT_DIR / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
