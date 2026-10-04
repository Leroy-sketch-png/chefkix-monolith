from __future__ import annotations

import copy
import json
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[2]
REFERENCE_DOCX = Path(r"C:\Users\LENOVO\Downloads\PhanPhuTho_VoTrungTin_ Sprint2.docx")
OUT_DIR = ROOT / "output" / "sprints"
FONT_DIR = ROOT / "tmp" / "source-inspect" / "fonts"
OUT_DIR.mkdir(parents=True, exist_ok=True)
FONT_DIR.mkdir(parents=True, exist_ok=True)


def sprint(
    n: int,
    date: str,
    summary: list[tuple[str, str]],
    objective: str,
    scope: str,
    done: list[tuple[str, str, str]],
    features: list[tuple[str, str, str, str]],
    evidence: list[str],
    risk: str,
    mitigation: str,
    next_backend: str,
    next_frontend: str,
    assignments: list[tuple[str, str, str]],
    outline_item: str,
    outline_status: str,
    outline_note: str,
    conclusion: str,
) -> dict:
    return locals()


REPORTS = [
    sprint(
        1,
        "10/02/2026",
        [
            ("Định hướng sản phẩm:", "Thống nhất bài toán xây dựng nền tảng chia sẻ ẩm thực có hướng dẫn nấu ăn từng bước và cơ chế gamification để tăng động lực học nấu."),
            ("Phân tích yêu cầu:", "Hoàn thiện các nhóm người dùng, user journey, use case chính và backlog phiên bản MVP cho luồng đăng ký, công thức, cooking session và tương tác xã hội."),
            ("Định hướng kỹ thuật:", "Đánh giá các lựa chọn Java Spring Boot, Next.js, MongoDB, Kafka, Redis và Keycloak; xác định các rủi ro kỹ thuật cần kiểm chứng ở các sprint sau."),
        ],
        "Chốt phạm vi MVP, làm rõ nhu cầu người dùng và tạo backlog đủ cụ thể để chuyển sang thiết kế hệ thống.",
        "Khảo sát bài toán, user journey, use case, phân tích công nghệ, product backlog và kế hoạch sprint.",
        [
            ("1.1", "Xác định vấn đề, đối tượng sử dụng và giá trị cốt lõi của nền tảng", "Done"),
            ("1.2", "Đặc tả user journey cho đăng nhập, công thức và cooking session", "Done"),
            ("1.3", "Lập product backlog và ưu tiên phạm vi MVP", "Done"),
            ("1.4", "Đánh giá stack kỹ thuật và các rủi ro ban đầu", "Done"),
        ],
        [
            ("Requirements", "MVP user journeys", "Done", "Đã thống nhất luồng chính cho người dùng mới và người dùng đã đăng nhập"),
            ("Product", "Backlog và acceptance criteria", "Done", "Các hạng mục được chia theo sprint để theo dõi"),
            ("Technical", "Technology evaluation", "Done", "Sẵn sàng chuyển sang thiết kế kiến trúc và dữ liệu"),
        ],
        [
            "Backlog và sơ đồ user journey được lưu trong thư mục minh chứng.",
            "Bảng đánh giá công nghệ và danh sách rủi ro kỹ thuật ban đầu.",
            "Biên bản thống nhất phạm vi MVP của nhóm.",
        ],
        "Phạm vi sản phẩm dễ mở rộng quá mức do có nhiều ý tưởng về AI, xã hội và gamification.",
        "Đóng khung MVP theo bốn luồng bắt buộc: Identity, Culinary, Social và Notification; các tính năng nâng cao được đưa vào backlog sau MVP.",
        "Thiết kế C4 Architecture, phân ranh giới module, ERD, API contract và phương án giao tiếp giữa các domain.",
        "Dựng prototype Figma cho cooking session, profile, feed và leaderboard; chuẩn bị component inventory.",
        [
            ("Võ Trung Tín", "Phân tích kiến trúc, công nghệ backend và rủi ro tích hợp", "Technical discovery"),
            ("Phan Phú Thọ", "User journey, backlog, wireframe và acceptance criteria", "Product and UX discovery"),
        ],
        "Khởi động và phân tích yêu cầu (26/01 – 10/02)",
        "Hoàn thành",
        "Phạm vi MVP và backlog được thống nhất, tạo cơ sở cho giai đoạn thiết kế hệ thống.",
        "Sprint 1 đã tạo được cùng một cách hiểu về sản phẩm và giới hạn phạm vi cần làm. Nhóm chuyển sang Sprint 2 với các đầu ra cụ thể hơn về kiến trúc, cơ sở dữ liệu, API và giao diện.",
    ),
    sprint(
        2,
        "23/02/2026",
        [
            ("Kiến trúc hệ thống:", "Đã chốt sơ đồ thiết kế Modular Monolith với 4 domain: Identity, Culinary, Social, Notification. Lên phương án giao tiếp nội bộ qua Java Interface và giao tiếp bất đồng bộ qua Kafka."),
            ("Database:", "Hoàn tất bản vẽ thiết kế schema dữ liệu (ERD), đặc biệt mở rộng để hỗ trợ gamification với UserStats, PointHistory và Badge."),
            ("UI/UX:", "Đã hoàn thiện prototype cho luồng Cooking Session (Timer, Step-by-step) và hệ thống Gamification (Profile Badge, Leaderboard)."),
        ],
        "Tạo nền móng kỹ thuật trên sơ đồ trước khi bắt tay vào code và bảo đảm UI/UX khả thi để triển khai bằng Next.js.",
        "System Architecture Diagram (C4 Model), Database ERD, UI/UX Prototype và API Contract.",
        [
            ("2.1", "Vẽ sơ đồ kiến trúc Modular Monolith (C4 Level 1 và 2)", "Done"),
            ("2.2", "Thiết kế API Contract sơ bộ (OpenAPI docs)", "Done"),
            ("2.3", "Thiết kế CSDL và ERD hỗ trợ Gamification", "Done"),
            ("2.4", "Thiết kế prototype giao diện trên Figma", "Done"),
        ],
        [
            ("Architecture", "Modular Monolith Diagram", "Done", "Đã chốt cấu trúc 4 module"),
            ("Database", "Gamification và Social Schema", "Done", "Sẵn sàng chuyển sang bước migrate"),
            ("UI/UX", "Cooking Session và Leaderboard", "Done", "Đã hoàn thiện mockup và luồng chính"),
        ],
        [
            "Prototype và hình ảnh sơ đồ được lưu trong thư mục minh chứng Google Drive.",
            "Sơ đồ C4, ERD và API contract được dùng làm tài liệu triển khai.",
            "Link thư mục minh chứng: https://drive.google.com/drive/folders/15QlKIiRh2G-62sHReu6Ui1QDdORpYUph",
        ],
        "Boundary giữa Culinary (công thức) và Social (comment, like) có thể chồng chéo dữ liệu, làm query phức tạp.",
        "Tuân thủ nguyên tắc modular; dùng Provider Interface cho gọi nội bộ và Kafka event cho các luồng bất đồng bộ thay vì query trực tiếp chéo module.",
        "Setup Docker Compose, Spring Boot multi-module, Keycloak, Kafka, Redis và khởi tạo Identity Module.",
        "Khởi tạo Next.js 15, Tailwind CSS, linting, formatter, Base Layout và các UI component cơ bản.",
        [
            ("Võ Trung Tín", "Thiết kế backend, API contract và phương án Docker, Keycloak, Kafka", "Backend Phase 1"),
            ("Phan Phú Thọ", "Prototype Figma, UI flow và component inventory", "Frontend and UX Phase 1"),
        ],
        "Giai đoạn Thiết kế Hệ thống và UI/UX (11/02 – 22/02)",
        "Hoàn thành",
        "Kiến trúc, database và Figma được hoàn thiện đúng timeline, sẵn sàng cho pha cài đặt.",
        "Sprint 2 đã chốt các quyết định thiết kế quan trọng để giảm việc làm lại khi code. Sprint 3 sẽ bắt đầu thiết lập môi trường và triển khai những thành phần đầu tiên của hệ thống.",
    ),
    sprint(
        3,
        "08/03/2026",
        [
            ("Backend:", "Khởi tạo Spring Boot 3 multi-module theo ranh giới Identity, Culinary, Social và Notification; chuẩn hóa shared kernel và Provider Interface."),
            ("Infrastructure:", "Dựng Docker Compose cho MongoDB, Kafka, Redis và Keycloak; bổ sung cấu hình môi trường và health check để nhóm có thể chạy stack nhất quán."),
            ("Frontend:", "Khởi tạo Next.js 15 với Tailwind CSS, TypeScript, ESLint, Prettier và Base Layout; kết nối thử frontend với API health endpoint."),
        ],
        "Đưa thiết kế Sprint 2 vào một bộ mã nguồn có thể build và chạy cục bộ, đồng thời thống nhất cách làm việc giữa hai thành viên.",
        "Project bootstrap, local infrastructure, security baseline, API health check, base layout và coding conventions.",
        [
            ("3.1", "Khởi tạo repository và Spring Boot multi-module", "Done"),
            ("3.2", "Dựng Docker Compose cho MongoDB, Kafka, Redis và Keycloak", "Done"),
            ("3.3", "Tích hợp Security baseline và kiểm tra JWT flow", "Done"),
            ("3.4", "Khởi tạo Next.js 15, Base Layout và UI components", "Done"),
        ],
        [
            ("Backend", "Multi-module build", "Done", "Các module build được trong cùng một Maven project"),
            ("Infrastructure", "Local service stack", "Done", "Có cấu hình khởi động và kiểm tra trạng thái"),
            ("Frontend", "Next.js foundation", "Done", "Có layout, style system và quy tắc linting"),
        ],
        [
            "Mã nguồn module, pom.xml, application configuration và Docker Compose.",
            "README hướng dẫn khởi động local stack và các endpoint kiểm tra trạng thái.",
            "Ảnh chụp màn hình frontend chạy được cùng backend development profile.",
        ],
        "Khác biệt phiên bản runtime hoặc biến môi trường có thể làm nhóm không khởi động được toàn bộ stack.",
        "Chuẩn hóa file .env.example, README quick start, profile development và health check cho từng phụ thuộc hạ tầng.",
        "Hoàn thiện Recipe CRUD, upload ảnh, tìm kiếm cơ bản và Cooking Session API.",
        "Dựng trang tạo công thức, recipe detail, bước nấu và timer state; bổ sung form validation.",
        [
            ("Võ Trung Tín", "Spring Boot modules, Docker Compose, Keycloak, Kafka và Redis", "Backend and infrastructure"),
            ("Phan Phú Thọ", "Next.js foundation, layout, form components và API client", "Frontend foundation"),
        ],
        "Giai đoạn Cài đặt nền tảng (23/02 – 08/03)",
        "Hoàn thành",
        "Môi trường, cấu trúc project và các quy ước code đã được cài đặt để phát triển tính năng.",
        "Sprint 3 đã chuyển thiết kế thành nền tảng chạy được ở local. Nhóm có thể bắt đầu phát triển các luồng nghiệp vụ chính thay vì chỉ kiểm thử cấu hình.",
    ),
    sprint(
        4,
        "22/03/2026",
        [
            ("Culinary core:", "Triển khai luồng tạo, cập nhật, xem và tìm kiếm công thức; mô hình hóa ingredients, steps, timers và metadata để phục vụ nhiều kiểu công thức."),
            ("Cooking Session:", "Hoàn thiện session theo từng bước, timer event và lưu tiến độ để người dùng có thể tiếp tục nấu sau khi rời trang."),
            ("Frontend:", "Kết nối recipe list, recipe detail, create recipe và cooking session với API; xử lý loading, error và empty state cơ bản."),
        ],
        "Hoàn thành luồng giá trị cốt lõi: người dùng tạo hoặc chọn một công thức và nấu theo từng bước có timer.",
        "Recipe domain, image upload contract, cooking session lifecycle, timer events, validation và các màn hình frontend tương ứng.",
        [
            ("4.1", "Triển khai Recipe CRUD và validation dữ liệu", "Done"),
            ("4.2", "Triển khai Cooking Session và timer event", "Done"),
            ("4.3", "Kết nối recipe list, detail và create flow", "Done"),
            ("4.4", "Bổ sung test cho service và controller chính", "Done"),
        ],
        [
            ("Recipe", "Create, edit, detail và search", "Done", "Các trường công thức và step được kiểm tra trước khi lưu"),
            ("Cooking", "Session lifecycle và timer", "Done", "Có thể bắt đầu, chuyển bước, pause và hoàn tất session"),
            ("Frontend", "Recipe and cooking screens", "Done", "Đã kết nối API và có trạng thái lỗi cơ bản"),
        ],
        [
            "Ảnh chụp recipe detail, create recipe và cooking session.",
            "Unit test cho RecipeService và CookingSessionService.",
            "API contract và mã nguồn trong các module culinary và frontend services.",
        ],
        "Recipe steps có cấu trúc linh hoạt, dễ tạo dữ liệu không nhất quán giữa form frontend và backend.",
        "Dùng DTO/schema validation ở cả hai phía, chuẩn hóa step number, ingredient reference và quy tắc timer trước khi lưu.",
        "Phát triển Social Feed, post, comment, like, save và các Provider Interface liên module.",
        "Dựng feed, post detail, comment composer và trạng thái like/save; chuẩn bị navigation giữa culinary và social.",
        [
            ("Võ Trung Tín", "Recipe, Cooking Session, timer events và service tests", "Culinary core"),
            ("Phan Phú Thọ", "Recipe screens, cooking UI, form validation và state handling", "Frontend culinary"),
        ],
        "Giai đoạn Phát triển Culinary Core (09/03 – 22/03)",
        "Hoàn thành",
        "Luồng công thức và cooking session đã có thể trình diễn đầu-cuối ở môi trường phát triển.",
        "Sprint 4 hoàn thiện hành trình nấu ăn cơ bản, là nền tảng để thêm tương tác xã hội và co-cooking trong các sprint tiếp theo.",
    ),
    sprint(
        5,
        "05/04/2026",
        [
            ("Social:", "Triển khai post, comment, like, save và feed; sử dụng PostProvider để giữ ranh giới giữa Social và các module khác."),
            ("Real-time:", "Dựng chat conversation, message và WebSocket/STOMP cho thông báo và tương tác theo thời gian thực trong nhóm nấu."),
            ("Frontend:", "Kết nối feed, post detail, chat và cooking room; đồng bộ state bằng store riêng để tránh làm gián đoạn cooking session."),
        ],
        "Mở rộng sản phẩm từ trải nghiệm cá nhân sang chia sẻ, thảo luận và nấu cùng người khác.",
        "Social feed, post/comment/like/save, chat, WebSocket/STOMP, cooking room và trạng thái realtime trên frontend.",
        [
            ("5.1", "Triển khai post, comment, like và save", "Done"),
            ("5.2", "Triển khai feed và PostProvider contract", "Done"),
            ("5.3", "Triển khai chat conversation, message và WebSocket", "Done"),
            ("5.4", "Kết nối feed và chat vào frontend", "Done"),
        ],
        [
            ("Social", "Feed và interaction", "Done", "Có post, comment, like, save và các trạng thái empty/error"),
            ("Realtime", "Chat và WebSocket", "Done", "Tin nhắn và notification channel có thể cập nhật theo thời gian thực"),
            ("Co-cooking", "Cooking room", "In progress", "Đã có service và UI flow, tiếp tục harden ở sprint sau"),
        ],
        [
            "Ảnh chụp feed, post detail, comment và chat.",
            "Test cho social service, chat flow và cooking room lifecycle.",
            "WebSocket configuration và các Provider Interface trong backend.",
        ],
        "Realtime state có thể bị lệch khi người dùng đổi trang hoặc mất kết nối ngắn hạn.",
        "Tách store theo feature, xử lý reconnect, hiển thị trạng thái kết nối và ưu tiên server state khi đồng bộ lại.",
        "Hoàn thiện XP, leaderboard, challenges, achievements và notification event consumers.",
        "Dựng profile progress, challenge cards, leaderboard, notification center và các trạng thái cập nhật realtime.",
        [
            ("Võ Trung Tín", "Social contracts, chat/WebSocket, cooking room và notification events", "Backend social and realtime"),
            ("Phan Phú Thọ", "Feed, post detail, chat, cooking room UI và state stores", "Frontend social and realtime"),
        ],
        "Giai đoạn Social và Realtime (23/03 – 05/04)",
        "Cơ bản hoàn thành",
        "Feed và chat đã hoạt động; cooking room cần tiếp tục kiểm thử reconnect và các tình huống nhiều người dùng.",
        "Sprint 5 đưa tính xã hội vào sản phẩm và tạo cơ sở cho co-cooking. Rủi ro chính hiện tại chuyển từ xây tính năng sang đồng bộ trạng thái và độ ổn định realtime.",
    ),
    sprint(
        6,
        "19/04/2026",
        [
            ("Gamification:", "Triển khai XP, statistics, streak, badge, challenge pool, duel và leaderboard để biến tiến trình nấu thành vòng lặp có động lực."),
            ("Notification:", "Kết nối event từ hoạt động người dùng tới bell notification, email và push token; bổ sung cleanup khi người dùng hoặc post bị xóa."),
            ("Safety baseline:", "Đưa kiểm tra nội dung, anti-cheat và rule-based moderation vào các điểm có thể ảnh hưởng đến XP hoặc nội dung cộng đồng."),
        ],
        "Hoàn thiện vòng lặp tiến bộ của người dùng và bảo đảm phần thưởng, thông báo và nội dung có thể kiểm soát được.",
        "XP flow, statistics, streak, badges, challenges, duels, leaderboard, notification consumers và moderation baseline.",
        [
            ("6.1", "Triển khai XP, statistics, streak và badge", "Done"),
            ("6.2", "Triển khai challenge lifecycle, pool và duel", "Done"),
            ("6.3", "Triển khai bell/email/push notification flow", "Done"),
            ("6.4", "Bổ sung anti-cheat và moderation rule baseline", "Done"),
        ],
        [
            ("Gamification", "XP, badge, streak và leaderboard", "Done", "Tiến trình người dùng có thể hiển thị và tính lại từ activity"),
            ("Challenges", "Challenge pool và duel", "Done", "Có lifecycle và các service test chính"),
            ("Notifications", "Bell, email và push", "Done", "Event consumer xử lý được các sự kiện chính"),
        ],
        [
            "Ảnh chụp profile progress, leaderboard, challenge và notification center.",
            "Test cho challenge lifecycle, statistics XP flow và notification cleanup.",
            "Rule set cho moderation và anti-cheat trong backend/AI service.",
        ],
        "Nếu XP được phát từ nhiều event mà không có idempotency, người dùng có thể nhận thưởng trùng.",
        "Gắn event identity, kiểm tra trạng thái trước khi cộng XP và bổ sung test cho retry, duplicate event và concurrent update.",
        "Tích hợp AI service, content moderation, allergen safety, recipe enrichment, search và graph endpoint.",
        "Kết nối AI actions vào create recipe, pantry/cooking assistant và các màn hình cần hiển thị safety status.",
        [
            ("Võ Trung Tín", "Gamification domain, challenge services, event consumers và idempotency", "Progress and notifications"),
            ("Phan Phú Thọ", "Profile progress, leaderboard, challenges và notification UI", "Frontend gamification"),
        ],
        "Giai đoạn Gamification và Trust Baseline (06/04 – 19/04)",
        "Hoàn thành",
        "Các vòng lặp XP, challenge, notification và các kiểm soát an toàn ban đầu đã được nối vào hệ thống.",
        "Sprint 6 hoàn thiện phần động lực và trust baseline của sản phẩm. Sprint 7 tập trung đưa AI, tìm kiếm và các hợp đồng dữ liệu nâng cao vào trải nghiệm end-to-end.",
    ),
    sprint(
        7,
        "03/05/2026",
        [
            ("AI service:", "Tích hợp FastAPI service cho recipe processing, enrichment, cooking assistant, moderation và các endpoint metadata; bổ sung provider rotation và graceful degradation."),
            ("Safety:", "Bảo vệ luồng thay thế nguyên liệu bằng allergen profile, giữ lại compound explanation và correlation id để theo dõi request."),
            ("Search and graph:", "Kết nối search/indexing và knowledge graph contract; frontend xử lý dữ liệu thật, trạng thái pending và lỗi tích hợp mà không bịa dữ liệu."),
        ],
        "Đưa các khả năng AI và dữ liệu nâng cao vào sản phẩm theo hướng có kiểm soát, có fallback và có thể quan sát.",
        "AI API integration, moderation, allergen safety, enrichment, rate limiting, correlation id, search và graph explorer contract.",
        [
            ("7.1", "Kết nối AI endpoints cho recipe, assistant và moderation", "Done"),
            ("7.2", "Bổ sung allergen safety, fallback và rate limit", "Done"),
            ("7.3", "Kết nối search và knowledge graph contract", "Done"),
            ("7.4", "Bổ sung test contract và observability cho các luồng AI", "Done"),
        ],
        [
            ("AI", "Recipe processing và enrichment", "Done", "Có response adapter và xử lý lỗi theo feature"),
            ("Trust", "Moderation và allergen safety", "Done", "Luồng nguy cơ cao fail closed hoặc chuyển sang review"),
            ("Data", "Search và knowledge graph", "Done", "Frontend dùng contract sống và hiển thị Pending khi thiếu dữ liệu"),
        ],
        [
            "API examples, health check và smoke tests của AI service.",
            "Test cho moderation, allergen safety, rate limiting, compound payload và correlation id.",
            "Tài liệu graph integration và ảnh chụp các trạng thái search/graph trên frontend.",
        ],
        "Phụ thuộc vào AI provider hoặc dữ liệu graph bên ngoài có thể làm trải nghiệm không ổn định và khó tái hiện lỗi.",
        "Provider rotation, cooldown, cache, timeout, graceful degradation, request id và trạng thái integration-pending rõ ràng trên UI.",
        "Chạy kiểm thử tích hợp toàn hệ thống, seed dữ liệu demo, rà soát security, performance và sửa lỗi trước nghiệm thu.",
        "Hoàn thiện demo flow, visual QA, empty/error states, accessibility check và tài liệu hướng dẫn trình diễn.",
        [
            ("Võ Trung Tín", "AI proxy, safety contract, observability, search và backend integration tests", "AI and platform integration"),
            ("Phan Phú Thọ", "Graph/search UI, AI states, visual QA và demo flow", "Frontend integration and QA"),
        ],
        "Giai đoạn Tích hợp AI và Dữ liệu nâng cao (20/04 – 03/05)",
        "Hoàn thành",
        "Các hợp đồng AI, safety, search và graph đã được tích hợp theo hướng có fallback và kiểm soát trạng thái.",
        "Sprint 7 mở rộng sản phẩm vượt ngoài CRUD nhưng vẫn giữ nguyên tắc không che giấu lỗi tích hợp. Sprint 8 sẽ tập trung vào độ ổn định, demo và bộ hồ sơ nghiệm thu.",
    ),
    sprint(
        8,
        "17/05/2026",
        [
            ("Integration:", "Chạy thử luồng end-to-end từ đăng nhập, tạo công thức, AI hỗ trợ, cooking session, social interaction tới XP và notification."),
            ("Quality:", "Rà soát test backend, AI service và frontend; bổ sung smoke test, runtime status, source fingerprint và các gate cho demo để giảm rủi ro trình diễn."),
            ("Delivery:", "Hoàn thiện Docker/Kubernetes configuration, seed dữ liệu demo, README quick start, checklist vận hành và tài liệu hướng dẫn bảo vệ đồ án."),
        ],
        "Đóng gói sản phẩm ở trạng thái có thể trình diễn và bàn giao, với bằng chứng kiểm thử và hướng dẫn khởi động rõ ràng.",
        "End-to-end demo, regression test, visual QA, seed data, runtime readiness, deployment configuration, documentation và handover.",
        [
            ("8.1", "Chạy regression và smoke test cho các luồng chính", "Done"),
            ("8.2", "Chuẩn bị seed dữ liệu, demo profile và readiness gate", "Done"),
            ("8.3", "Rà soát cấu hình Docker/Kubernetes và source fingerprint", "Done"),
            ("8.4", "Hoàn thiện tài liệu triển khai, demo và bàn giao", "Done"),
        ],
        [
            ("End-to-end", "Core user journeys", "Done", "Có thể trình diễn từ identity tới cooking và social"),
            ("Quality", "Automated checks và smoke matrix", "Done", "Có test report và gate để phát hiện lỗi trước demo"),
            ("Delivery", "Runtime và handover", "Done", "Có compose, seed, README và checklist vận hành"),
        ],
        [
            "Bộ test report, smoke test và runtime readiness trong các repository.",
            "Ảnh chụp demo, dashboard trạng thái và các màn hình chính của sản phẩm.",
            "Link thư mục minh chứng: https://drive.google.com/drive/folders/15QlKIiRh2G-62sHReu6Ui1QDdORpYUph",
        ],
        "Khác biệt giữa môi trường máy cá nhân và môi trường trình diễn có thể làm thay đổi trạng thái dịch vụ hoặc dữ liệu.",
        "Dùng seed dữ liệu xác định, checklist readiness, source fingerprint, health/status checks và một kịch bản demo có đường lui khi AI hoặc dịch vụ phụ trợ không sẵn sàng.",
        "Theo dõi sau nghiệm thu, xử lý các lỗi còn lại từ demo, bổ sung tài liệu và chuẩn bị kế hoạch mở rộng sau đồ án.",
        "Đóng băng giao diện trình diễn, chuẩn bị slide/screenshot, rehearsal bảo vệ và rà soát các yêu cầu còn thiếu.",
        [
            ("Võ Trung Tín", "Regression backend/AI, infrastructure readiness, seed và demo runtime", "Release and platform"),
            ("Phan Phú Thọ", "Frontend visual QA, demo walkthrough, screenshots và presentation assets", "Release and presentation"),
        ],
        "Giai đoạn Tích hợp, Kiểm thử và Bàn giao (04/05 – 17/05)",
        "Hoàn thành",
        "Hệ thống và hồ sơ bàn giao đạt trạng thái sẵn sàng để trình diễn; các giới hạn phụ thuộc môi trường đã được ghi nhận.",
        "Sau 8 sprint, nhóm đã đi từ phân tích yêu cầu tới một nền tảng cooking social có modular backend, frontend, AI service, gamification, realtime và hạ tầng chạy được. Các đầu ra quan trọng đều có mã nguồn, test hoặc tài liệu đi kèm để phục vụ báo cáo và nghiệm thu.",
    ),
]


def set_run_font(run, size: float = 12, bold: bool | None = None, italic: bool | None = None):
    run.font.name = "Google Sans"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Google Sans")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Google Sans")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def fill_paragraph(p, segments, size=12):
    clear_paragraph(p)
    for text, bold, italic in segments:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic)


def set_heading(p, text, level):
    p.style = {2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}[level]
    fill_paragraph(p, [(text, True, level == 4)], size={2: 18, 3: 14, 4: 12}[level])


def set_cell(cell, text, bold=False, size=11.5):
    p = cell.paragraphs[0]
    fill_paragraph(p, [(text, bold, False)], size=size)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def fill_docx(report: dict, out_path: Path):
    doc = Document(str(REFERENCE_DOCX))
    p = doc.paragraphs
    set_heading(p[0], f"BÁO CÁO TIẾN ĐỘ 2 TUẦN (SPRINT {report['n']})", 2)
    fill_paragraph(p[1], [("Đồ án:", True, False), (" Step-by-Step Cooking Platform – Nền tảng Chia sẻ Ẩm thực Gamified", False, False)])
    fill_paragraph(p[2], [("Trường:", True, False), (" Trường Đại học Công nghệ Thông tin - ĐHQG TP.HCM", False, False)])
    fill_paragraph(p[3], [("Giảng viên hướng dẫn:", True, False), (" ThS. Trần Thị Hồng Yến", False, False)])
    fill_paragraph(p[4], [("Nhóm thực hiện:", True, False), (" Phan Phú Thọ, Võ Trung Tín", False, False)])
    fill_paragraph(p[5], [("Ngày báo cáo:", True, False), (f" {report['date']}", False, False)])
    set_heading(p[6], "1. TÓM TẮT ĐIỀU HÀNH", 3)
    for idx, (label, text) in enumerate(report["summary"], start=7):
        fill_paragraph(p[idx], [(label, True, False), (f" {text}", False, False)])
    set_heading(p[10], "2. MỤC TIÊU & PHẠM VI", 3)
    fill_paragraph(p[11], [("Mục tiêu:", True, False), (f" {report['objective']}", False, False)])
    fill_paragraph(p[12], [("Phạm vi:", True, False), (f" {report['scope']}", False, False)])
    set_heading(p[13], "3. CÁC HẠNG MỤC ĐÃ HOÀN THÀNH", 3)
    set_heading(p[14], "4. BẢNG TRẠNG THÁI TÍNH NĂNG", 3)
    set_heading(p[15], "5. KẾT QUẢ MINH CHỨNG", 3)
    for idx, text in enumerate(report["evidence"], start=16):
        fill_paragraph(p[idx], [(text, False, False)])
    set_heading(p[19], "6. RỦI RO & BIỆN PHÁP", 3)
    set_heading(p[20], f"7. KẾ HOẠCH 2 TUẦN TỚI (Sprint {report['n'] + 1})", 3)
    set_heading(p[21], "7.1 Backend và Infrastructure", 4)
    fill_paragraph(p[22], [(report["next_backend"], False, False)])
    set_heading(p[23], "7.2 Frontend và QA", 4)
    fill_paragraph(p[24], [(report["next_frontend"], False, False)])
    set_heading(p[25], f"8. PHÂN CÔNG (Sprint {report['n'] + 1})", 3)
    set_heading(p[26], "9. ĐỐI CHIẾU VỚI ĐỀ CƯƠNG", 3)
    set_heading(p[27], "10. KẾT LUẬN", 3)
    fill_paragraph(p[28], [(report["conclusion"], False, False)])
    set_heading(p[29], "11. Phụ lục A - Liên kết và Minh hoạ", 3)
    fill_paragraph(p[30], [("Link minh chứng:", True, False), (" https://drive.google.com/drive/folders/15QlKIiRh2G-62sHReu6Ui1QDdORpYUph", False, False)])
    fill_paragraph(p[31], [("Người lập báo cáo:", True, False), (" Võ Trung Tín", False, False)])

    tables = doc.tables
    headers = [
        ["Mục", "Nội dung", "Trạng thái"],
        ["Nhóm chức năng", "Tiểu mục", "Trạng thái", "Ghi chú"],
        ["Rủi ro", "Biện pháp khắc phục"],
        ["Thành viên", "Phụ trách", "Ghi chú"],
        ["Hạng mục trong Đề cương", "Trạng thái thực tế", "Ghi chú/Mức độ hoàn thành"],
    ]
    for ti, row in enumerate(headers):
        for ci, text in enumerate(row):
            set_cell(tables[ti].cell(0, ci), text, bold=True)
    for ri, row in enumerate(report["done"], start=1):
        for ci, text in enumerate(row):
            set_cell(tables[0].cell(ri, ci), text, bold=False)
    for ri, row in enumerate(report["features"], start=1):
        for ci, text in enumerate(row):
            set_cell(tables[1].cell(ri, ci), text, bold=False)
    set_cell(tables[2].cell(1, 0), report["risk"])
    set_cell(tables[2].cell(1, 1), report["mitigation"])
    for ri, row in enumerate(report["assignments"], start=1):
        for ci, text in enumerate(row):
            set_cell(tables[3].cell(ri, ci), text, bold=(ci == 0))
    outline_row = [report["outline_item"], report["outline_status"], report["outline_note"]]
    for ci, text in enumerate(outline_row):
        set_cell(tables[4].cell(1, ci), text, bold=(ci == 1))
    doc.core_properties.title = f"Báo cáo tiến độ Sprint {report['n']} - ChefKix"
    doc.core_properties.subject = "Báo cáo tiến độ đồ án Step-by-Step Cooking Platform"
    doc.core_properties.author = "Võ Trung Tín"
    doc.save(str(out_path))


def extract_fonts():
    with zipfile.ZipFile(REFERENCE_DOCX) as z:
        for name, target in [
            ("word/fonts/GoogleSans-regular.ttf", FONT_DIR / "GoogleSans-regular.ttf"),
            ("word/fonts/GoogleSans-bold.ttf", FONT_DIR / "GoogleSans-bold.ttf"),
            ("word/fonts/GoogleSans-italic.ttf", FONT_DIR / "GoogleSans-italic.ttf"),
        ]:
            target.write_bytes(z.read(name))


def pdf_font_setup():
    pdfmetrics.registerFont(TTFont("GoogleSans", str(FONT_DIR / "GoogleSans-regular.ttf")))
    pdfmetrics.registerFont(TTFont("GoogleSans-Bold", str(FONT_DIR / "GoogleSans-bold.ttf")))
    pdfmetrics.registerFont(TTFont("GoogleSans-Italic", str(FONT_DIR / "GoogleSans-italic.ttf")))


def pdf_styles():
    return {
        "body": ParagraphStyle("body", fontName="GoogleSans", fontSize=11.2, leading=14.2, textColor=colors.HexColor("#1F1F1F"), spaceAfter=8),
        "title": ParagraphStyle("title", fontName="GoogleSans-Bold", fontSize=18, leading=21, textColor=colors.HexColor("#1F1F1F"), spaceAfter=8),
        "h3": ParagraphStyle("h3", fontName="GoogleSans-Bold", fontSize=14, leading=17, textColor=colors.HexColor("#1F1F1F"), spaceBefore=7, spaceAfter=7),
        "h4": ParagraphStyle("h4", fontName="GoogleSans-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#1F1F1F"), spaceBefore=4, spaceAfter=6),
        "table": ParagraphStyle("table", fontName="GoogleSans", fontSize=10, leading=12.2, textColor=colors.HexColor("#1F1F1F")),
        "table_bold": ParagraphStyle("table_bold", fontName="GoogleSans-Bold", fontSize=10, leading=12.2, textColor=colors.HexColor("#1F1F1F")),
    }


def esc(text: str) -> str:
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def ptext(text: str, style, bold_label: str | None = None):
    if bold_label and text.startswith(bold_label):
        return Paragraph(f"<b>{esc(bold_label)}</b>{esc(text[len(bold_label):])}", style)
    return Paragraph(esc(text), style)


def table_flow(data, styles, widths, bold_rows=(0,), bold_cells=None):
    bold_cells = bold_cells or set()
    rows = []
    for ri, row in enumerate(data):
        rendered = []
        for ci, value in enumerate(row):
            use_bold = ri in bold_rows or (ri, ci) in bold_cells
            rendered.append(Paragraph(esc(str(value)), styles["table_bold" if use_bold else "table"]))
        rows.append(rendered)
    t = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#1F1F1F")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F4F8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]
    for ri in range(1, len(rows)):
        if ri % 2 == 0:
            commands.append(("BACKGROUND", (0, ri), (-1, ri), colors.HexColor("#FAFBFC")))
    t.setStyle(TableStyle(commands))
    return t


def build_pdf(report: dict, out_path: Path):
    styles = pdf_styles()
    doc = BaseDocTemplate(
        str(out_path),
        pagesize=A4,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title=f"Báo cáo tiến độ Sprint {report['n']} - ChefKix",
        author="Võ Trung Tín",
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame])])
    story = []
    story.append(Paragraph(f"BÁO CÁO TIẾN ĐỘ 2 TUẦN (SPRINT {report['n']})", styles["title"]))
    metadata = [
        ("Đồ án:", "Step-by-Step Cooking Platform – Nền tảng Chia sẻ Ẩm thực Gamified"),
        ("Trường:", "Trường Đại học Công nghệ Thông tin - ĐHQG TP.HCM"),
        ("Giảng viên hướng dẫn:", "ThS. Trần Thị Hồng Yến"),
        ("Nhóm thực hiện:", "Phan Phú Thọ, Võ Trung Tín"),
        ("Ngày báo cáo:", report["date"]),
    ]
    for label, value in metadata:
        story.append(Paragraph(f"<b>{esc(label)}</b> {esc(value)}", styles["body"]))
    story.append(Paragraph("1. TÓM TẮT ĐIỀU HÀNH", styles["h3"]))
    for label, value in report["summary"]:
        story.append(Paragraph(f"<b>{esc(label)}</b> {esc(value)}", styles["body"]))
    story.append(Paragraph("2. MỤC TIÊU &amp; PHẠM VI", styles["h3"]))
    story.append(Paragraph(f"• <b>Mục tiêu:</b> {esc(report['objective'])}", styles["body"]))
    story.append(Paragraph(f"• <b>Phạm vi:</b> {esc(report['scope'])}", styles["body"]))
    story.append(Paragraph("3. CÁC HẠNG MỤC ĐÃ HOÀN THÀNH", styles["h3"]))
    story.append(table_flow([["Mục", "Nội dung", "Trạng thái"], *report["done"]], styles, [0.55 * inch, 3.35 * inch, 1.6 * inch]))
    story.append(Spacer(1, 10))
    story.append(Paragraph("4. BẢNG TRẠNG THÁI TÍNH NĂNG", styles["h3"]))
    story.append(table_flow([["Nhóm chức năng", "Tiểu mục", "Trạng thái", "Ghi chú"], *report["features"]], styles, [1.25 * inch, 1.7 * inch, 0.95 * inch, 2.1 * inch]))
    story.append(Spacer(1, 10))
    story.append(Paragraph("5. KẾT QUẢ MINH CHỨNG", styles["h3"]))
    for value in report["evidence"]:
        story.append(Paragraph(f"• {esc(value)}", styles["body"]))
    story.append(Paragraph("6. RỦI RO &amp; BIỆN PHÁP", styles["h3"]))
    story.append(table_flow([["Rủi ro", "Biện pháp khắc phục"], [report["risk"], report["mitigation"]]], styles, [3.25 * inch, 3.25 * inch]))
    story.append(Spacer(1, 10))
    story.append(Paragraph(f"7. KẾ HOẠCH 2 TUẦN TỚI (Sprint {report['n'] + 1})", styles["h3"]))
    story.append(Paragraph("7.1 Backend và Infrastructure", styles["h4"]))
    story.append(Paragraph(esc(report["next_backend"]), styles["body"]))
    story.append(Paragraph("7.2 Frontend và QA", styles["h4"]))
    story.append(Paragraph(esc(report["next_frontend"]), styles["body"]))
    story.append(Paragraph(f"8. PHÂN CÔNG (Sprint {report['n'] + 1})", styles["h3"]))
    story.append(table_flow([["Thành viên", "Phụ trách", "Ghi chú"], *report["assignments"]], styles, [1.7 * inch, 3.35 * inch, 1.45 * inch], bold_cells={(1, 0), (2, 0)}))
    story.append(Spacer(1, 10))
    story.append(Paragraph("9. ĐỐI CHIẾU VỚI ĐỀ CƯƠNG", styles["h3"]))
    story.append(table_flow([["Hạng mục trong Đề cương", "Trạng thái thực tế", "Ghi chú/Mức độ hoàn thành"], [report["outline_item"], report["outline_status"], report["outline_note"]],], styles, [2.35 * inch, 1.55 * inch, 2.6 * inch], bold_cells={(1, 1)}))
    story.append(Spacer(1, 10))
    story.append(Paragraph("10. KẾT LUẬN", styles["h3"]))
    story.append(Paragraph(esc(report["conclusion"]), styles["body"]))
    story.append(Paragraph("11. Phụ lục A - Liên kết và Minh hoạ", styles["h3"]))
    story.append(Paragraph("<b>Link minh chứng:</b> https://drive.google.com/drive/folders/15QlKIiRh2G-62sHReu6Ui1QDdORpYUph", styles["body"]))
    story.append(Paragraph("<b>Người lập báo cáo:</b> Võ Trung Tín", styles["body"]))
    doc.build(story)


def main():
    extract_fonts()
    pdf_font_setup()
    manifest = []
    for report in REPORTS:
        stem = f"PhanPhuTho_VoTrungTin_Sprint{report['n']}"
        docx_path = OUT_DIR / f"{stem}.docx"
        pdf_path = OUT_DIR / f"{stem}.pdf"
        fill_docx(report, docx_path)
        build_pdf(report, pdf_path)
        manifest.append({"sprint": report["n"], "docx": str(docx_path), "pdf": str(pdf_path)})
    (OUT_DIR / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
